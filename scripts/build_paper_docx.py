#!/usr/bin/env python3
"""Build a Chemical Industry and Engineering Progress manuscript DOCX.

The renderer consumes a structured YAML/JSON manuscript and applies direct
formatting based on STYLE_MAP.json. It deliberately avoids Word features that
are disallowed for a clean submission copy: fields, hyperlinks, text boxes,
floating images, footnotes, endnotes, comments, and section-heavy layout.
"""
from __future__ import annotations

import argparse
import html
import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

try:
    from PIL import Image, ImageDraw
except Exception:  # pragma: no cover - dependency is declared, but keep CLI usable.
    Image = None
    ImageDraw = None


DEFAULT_ZH_FONT = "宋体"
DEFAULT_LATIN_FONT = "Times New Roman"
EMU_PER_MM = 36_000
TWIPS_PER_CM = 567


def load_structured(path: Path) -> dict[str, Any]:
    if path.suffix.lower() in {".yaml", ".yml"}:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    else:
        data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise TypeError("manuscript source must be a YAML/JSON object")
    return data


def load_style(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def style_at(style: dict[str, Any], *path: str, default: Any = None) -> Any:
    cur: Any = style
    for key in path:
        if not isinstance(cur, dict) or key not in cur:
            return default
        cur = cur[key]
    return cur


def first_number(style: dict[str, Any], style_name: str, *keys: str, default: float) -> float:
    st = style_at(style, "paragraph_styles", style_name, default={})
    if not isinstance(st, dict):
        return default
    for key in keys:
        value = st.get(key)
        if isinstance(value, (int, float)):
            return float(value)
    return default


def clean_font_name(name: str | None, fallback: str) -> str:
    if not name:
        return fallback
    return str(name).split(";")[0].strip() or fallback


def allowed_zh_font(style: dict[str, Any], preferred: str | None, fallback: str = DEFAULT_ZH_FONT) -> str:
    allowed = set(style_at(style, "global_fonts", "chinese_allowed", default=[]) or [])
    font = clean_font_name(preferred, fallback)
    return font if not allowed or font in allowed else fallback


def get_or_add(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def remove_child(parent, tag: str) -> None:
    child = parent.find(qn(tag))
    if child is not None:
        parent.remove(child)


def strict_style(style: dict[str, Any], name: str) -> dict[str, Any]:
    value = style_at(style, "strict_layout", "styles", name, default={})
    return value if isinstance(value, dict) else {}


def strict_size_pt(style: dict[str, Any], name: str, default: float) -> float:
    value = strict_style(style, name).get("sz")
    if value is None:
        return default
    return float(value) / 2


def strict_font_zh(style: dict[str, Any], name: str, default: str = DEFAULT_ZH_FONT) -> str:
    return clean_font_name(strict_style(style, name).get("font_zh"), default)


def strict_font_latin(style: dict[str, Any], name: str, default: str = DEFAULT_LATIN_FONT) -> str:
    return clean_font_name(strict_style(style, name).get("font_latin"), default)


def strict_char_spacing(style: dict[str, Any], name: str) -> str | None:
    value = strict_style(style, name).get("character_spacing")
    return str(value) if value is not None else None


def alignment_from_ooxml(value: str | None) -> WD_ALIGN_PARAGRAPH | None:
    return {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }.get(str(value)) if value is not None else None


def set_raw_paragraph_attrs(paragraph, style: dict[str, Any], name: str) -> None:
    st = strict_style(style, name)
    defaults = style_at(style, "strict_layout", "paragraph_defaults", default={}) or {}
    p_pr = paragraph._p.get_or_add_pPr()

    jc = st.get("jc")
    if jc is not None:
        jc_el = get_or_add(p_pr, "w:jc")
        jc_el.set(qn("w:val"), str(jc))

    spacing_values = {
        "line": st.get("line", defaults.get("line")),
        "lineRule": st.get("lineRule", defaults.get("lineRule")),
        "before": st.get("before", defaults.get("before")),
        "after": st.get("after", defaults.get("after")),
    }
    spacing = get_or_add(p_pr, "w:spacing")
    for attr, value in spacing_values.items():
        if value is not None:
            spacing.set(qn(f"w:{attr}"), str(value))

    ind_values: dict[str, Any] = {}
    if isinstance(st.get("ind"), dict):
        ind_values.update(st["ind"])
    for key in ("firstLine", "hanging", "left", "right"):
        if st.get(key) is not None:
            ind_values[key] = st[key]
    if ind_values:
        ind = get_or_add(p_pr, "w:ind")
        for attr, value in ind_values.items():
            ind.set(qn(f"w:{attr}"), str(value))


def set_rfonts(run, zh: str = DEFAULT_ZH_FONT, latin: str = DEFAULT_LATIN_FONT) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:cs"), latin)
    r_fonts.set(qn("w:eastAsia"), zh)


def set_run_font(
    run,
    *,
    size_pt: float | None = None,
    zh: str = DEFAULT_ZH_FONT,
    latin: str = DEFAULT_LATIN_FONT,
    bold: bool | None = None,
    italic: bool | None = None,
    superscript: bool | None = None,
    sz_cs_pt: float | None = None,
    character_spacing: str | int | None = None,
) -> None:
    set_rfonts(run, zh=zh, latin=latin)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
        r_pr = run._element.get_or_add_rPr()
        sz_cs = get_or_add(r_pr, "w:szCs")
        sz_cs.set(qn("w:val"), str(int(round((sz_cs_pt if sz_cs_pt is not None else size_pt) * 2))))
    if character_spacing is not None:
        r_pr = run._element.get_or_add_rPr()
        spacing = get_or_add(r_pr, "w:spacing")
        spacing.set(qn("w:val"), str(character_spacing))
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if superscript is not None:
        run.font.superscript = superscript


def set_paragraph_format(
    paragraph,
    *,
    align: WD_ALIGN_PARAGRAPH | None = None,
    first_indent_cm: float | None = None,
    left_indent_cm: float | None = None,
    right_indent_cm: float | None = None,
    line_spacing: float = 1.0,
    space_before_pt: float = 0,
    space_after_pt: float = 0,
    keep_with_next: bool | None = None,
) -> None:
    if align is not None:
        paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_before = Pt(space_before_pt)
    pf.space_after = Pt(space_after_pt)
    if first_indent_cm is not None:
        pf.first_line_indent = Cm(first_indent_cm)
    if left_indent_cm is not None:
        pf.left_indent = Cm(left_indent_cm)
    if right_indent_cm is not None:
        pf.right_indent = Cm(right_indent_cm)
    if keep_with_next is not None:
        pf.keep_with_next = keep_with_next


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    size_pt: float = 10.5,
    zh: str = DEFAULT_ZH_FONT,
    latin: str = DEFAULT_LATIN_FONT,
    align: WD_ALIGN_PARAGRAPH | None = None,
    first_indent_cm: float | None = None,
    left_indent_cm: float | None = None,
    right_indent_cm: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    keep_with_next: bool | None = None,
    character_spacing: str | int | None = None,
    sz_cs_pt: float | None = None,
) -> Any:
    paragraph = doc.add_paragraph()
    set_paragraph_format(
        paragraph,
        align=align,
        first_indent_cm=first_indent_cm,
        left_indent_cm=left_indent_cm,
        right_indent_cm=right_indent_cm,
        keep_with_next=keep_with_next,
    )
    if text:
        run = paragraph.add_run(text)
        set_run_font(
            run,
            size_pt=size_pt,
            zh=zh,
            latin=latin,
            bold=bold,
            italic=italic,
            sz_cs_pt=sz_cs_pt,
            character_spacing=character_spacing,
        )
    return paragraph


def add_runs_from_text(
    paragraph,
    text: str,
    *,
    size_pt: float,
    zh: str,
    latin: str,
    bold: bool | None = None,
    sz_cs_pt: float | None = None,
    character_spacing: str | int | None = None,
) -> None:
    parts = str(text).split("\n")
    for i, part in enumerate(parts):
        if i:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        run = paragraph.add_run(part)
        set_run_font(run, size_pt=size_pt, zh=zh, latin=latin, bold=bold, sz_cs_pt=sz_cs_pt, character_spacing=character_spacing)


def add_labeled_paragraph(
    doc: Document,
    label: str,
    text: str,
    *,
    size_pt: float,
    zh: str,
    latin: str,
    align: WD_ALIGN_PARAGRAPH | None = None,
    left_indent_cm: float | None = None,
    right_indent_cm: float | None = None,
    label_bold: bool | None = None,
    character_spacing: str | int | None = None,
    sz_cs_pt: float | None = None,
) -> Any:
    paragraph = doc.add_paragraph()
    set_paragraph_format(
        paragraph,
        align=align,
        left_indent_cm=left_indent_cm,
        right_indent_cm=right_indent_cm,
    )
    label_run = paragraph.add_run(label)
    set_run_font(
        label_run,
        size_pt=size_pt,
        zh=zh,
        latin=latin,
        bold=label_bold,
        sz_cs_pt=sz_cs_pt,
        character_spacing=character_spacing,
    )
    body_run = paragraph.add_run(text)
    set_run_font(body_run, size_pt=size_pt, zh=zh, latin=latin, bold=None, sz_cs_pt=sz_cs_pt, character_spacing=character_spacing)
    return paragraph


def add_strict_paragraph(
    doc: Document,
    text: str,
    style: dict[str, Any],
    name: str,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    keep_with_next: bool | None = None,
) -> Any:
    st = strict_style(style, name)
    paragraph = doc.add_paragraph()
    set_raw_paragraph_attrs(paragraph, style, name)
    if keep_with_next is not None:
        paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(text)
    set_run_font(
        run,
        size_pt=strict_size_pt(style, name, 10.5),
        sz_cs_pt=float(st.get("szCs", st.get("sz", 21))) / 2,
        zh=strict_font_zh(style, name),
        latin=strict_font_latin(style, name),
        bold=bold if bold is not None else st.get("bold"),
        italic=italic,
        character_spacing=strict_char_spacing(style, name),
    )
    return paragraph


def set_run_from_strict_config(
    run,
    style: dict[str, Any],
    name: str,
    role: str | None = None,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    st = strict_style(style, name)
    cfg = st.get(role) if role else st
    if not isinstance(cfg, dict):
        cfg = st
    sz = cfg.get("sz", st.get("sz", 21))
    sz_cs = cfg.get("szCs", cfg.get("sz", st.get("szCs", st.get("sz", 21))))
    set_run_font(
        run,
        size_pt=float(sz) / 2,
        sz_cs_pt=float(sz_cs) / 2,
        zh=str(cfg.get("font_zh", st.get("font_zh", DEFAULT_ZH_FONT))),
        latin=str(cfg.get("font_latin", st.get("font_latin", DEFAULT_LATIN_FONT))),
        bold=bold if bold is not None else cfg.get("bold"),
        italic=italic,
        character_spacing=cfg.get("character_spacing", st.get("character_spacing")),
    )


def add_strict_labeled_paragraph(
    doc: Document,
    label: str,
    text: str,
    style: dict[str, Any],
    name: str,
    *,
    label_bold: bool | None = None,
) -> Any:
    st = strict_style(style, name)
    paragraph = doc.add_paragraph()
    set_raw_paragraph_attrs(paragraph, style, name)
    label_run = paragraph.add_run(label)
    set_run_from_strict_config(
        label_run,
        style,
        name,
        "label_run",
        bold=label_bold if label_bold is not None else st.get("label_bold"),
    )
    body_run = paragraph.add_run(text)
    set_run_from_strict_config(body_run, style, name, "body_run")
    return paragraph


def add_strict_classification_line(doc: Document, classification: dict[str, Any], style: dict[str, Any]) -> Any:
    paragraph = doc.add_paragraph()
    set_raw_paragraph_attrs(paragraph, style, "classification_line")

    def add_piece(text: str, role: str) -> None:
        run = paragraph.add_run(text)
        set_run_from_strict_config(run, style, "classification_line", role)

    add_piece("中图分类号：", "label_run")
    add_piece(str(classification.get("clc", "")), "body_run")
    add_piece("          ", "body_run")
    add_piece("文献标志码：", "label_run")
    add_piece(str(classification.get("document_code", "A")), "body_run")
    add_piece("      ", "body_run")
    add_piece("文章编号：", "label_run")
    add_piece(str(classification.get("article_number", "")), "body_run")
    return paragraph


def set_document_defaults(doc: Document, style: dict[str, Any]) -> None:
    normal = doc.styles["Normal"]
    strict_normal = style_at(style, "strict_layout", "normal_style", default={}) or {}
    body_size = float(strict_normal.get("sz", 21)) / 2
    normal.font.name = DEFAULT_LATIN_FONT
    normal.font.size = Pt(body_size)
    r_pr = normal._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    fonts = strict_normal.get("rFonts", {})
    r_fonts.set(qn("w:eastAsia"), clean_font_name(fonts.get("eastAsia"), DEFAULT_ZH_FONT))
    r_fonts.set(qn("w:ascii"), clean_font_name(fonts.get("ascii"), DEFAULT_LATIN_FONT))
    r_fonts.set(qn("w:hAnsi"), clean_font_name(fonts.get("hAnsi"), DEFAULT_LATIN_FONT))
    r_fonts.set(qn("w:cs"), clean_font_name(fonts.get("cs"), DEFAULT_LATIN_FONT))
    sz_cs = get_or_add(r_pr, "w:szCs")
    sz_cs.set(qn("w:val"), str(strict_normal.get("szCs", "24")))
    p_pr = normal._element.get_or_add_pPr()
    jc = get_or_add(p_pr, "w:jc")
    jc.set(qn("w:val"), str(strict_normal.get("jc", "both")))
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)


def apply_page_setup(doc: Document, style: dict[str, Any]) -> None:
    page = style.get("page_setup", {})
    section = doc.sections[0]
    section.page_width = Cm(page.get("width_cm", 21.0))
    section.page_height = Cm(page.get("height_cm", 29.7))
    section.top_margin = Cm(page.get("top_margin_cm", 2.5))
    section.bottom_margin = Cm(page.get("bottom_margin_cm", 2.0))
    section.left_margin = Cm(page.get("left_margin_cm", 2.0))
    section.right_margin = Cm(page.get("right_margin_cm", 2.0))
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)

    sect_pr = section._sectPr
    strict_page = style_at(style, "strict_layout", "page_ooxml", default={}) or {}
    pg_sz = get_or_add(sect_pr, "w:pgSz")
    for attr, value in (strict_page.get("pgSz") or {}).items():
        pg_sz.set(qn(f"w:{attr}"), str(value))
    pg_mar = get_or_add(sect_pr, "w:pgMar")
    for attr, value in (strict_page.get("pgMar") or {}).items():
        pg_mar.set(qn(f"w:{attr}"), str(value))
    cols = get_or_add(sect_pr, "w:cols")
    cols.set(qn("w:num"), str(page.get("columns", 1)))

    doc_grid = get_or_add(sect_pr, "w:docGrid")
    grid = strict_page.get("docGrid") or page.get("doc_grid", {})
    doc_grid.set(qn("w:type"), grid.get("type", "linesAndChars"))
    doc_grid.set(qn("w:linePitch"), str(grid.get("linePitch", 324)))
    doc_grid.set(qn("w:charSpace"), str(grid.get("charSpace", 4294967090)))


def build_author_paragraph(doc: Document, authors: list[dict[str, Any]], *, english: bool, style: dict[str, Any]) -> None:
    key = "en_authors" if english else "zh_authors"
    st = strict_style(style, key)
    size = strict_size_pt(style, key, 10.5)
    sz_cs_pt = float(st.get("szCs", st.get("sz", 21))) / 2
    zh = strict_font_zh(style, key)
    latin = strict_font_latin(style, key)
    separator = ", " if english else "，"
    paragraph = doc.add_paragraph()
    set_raw_paragraph_attrs(paragraph, style, key)
    for idx, author in enumerate(authors):
        if idx:
            run = paragraph.add_run(separator)
            set_run_font(run, size_pt=size, sz_cs_pt=sz_cs_pt, zh=zh, latin=latin)
        name = author.get("name_en" if english else "name_zh", "")
        name_run = paragraph.add_run(str(name))
        set_run_font(name_run, size_pt=size, sz_cs_pt=sz_cs_pt, zh=zh, latin=latin, italic=True if english else None)
        affiliations = author.get("affiliations", [])
        if affiliations:
            aff = ",".join(str(x) for x in affiliations)
            aff_run = paragraph.add_run(aff)
            set_run_font(aff_run, size_pt=size, sz_cs_pt=sz_cs_pt, zh=zh, latin=latin, superscript=True)


def affiliation_line(affiliations: list[dict[str, Any]], *, english: bool) -> str:
    if not affiliations:
        return ""
    if english:
        return "(" + "; ".join(f"{a.get('id', '')}{a.get('en', '')}" for a in affiliations) + ")"
    return "（" + "；".join(f"{a.get('id', '')}{a.get('zh', '')}" for a in affiliations) + "）"


class CitationManager:
    def __init__(self, references: list[dict[str, Any]]) -> None:
        self.references = references
        self.by_id = {str(ref.get("id")): ref for ref in references if ref.get("id") is not None}
        self.order: list[str] = []
        self.id_to_num: dict[str, int] = {}

    def _assign(self, ref_id: str) -> int:
        if ref_id not in self.id_to_num:
            self.order.append(ref_id)
            self.id_to_num[ref_id] = len(self.order)
        return self.id_to_num[ref_id]

    def cite_text(self, raw: str) -> str:
        ids = [x.strip() for x in re.split(r"[,;，、\s]+", raw) if x.strip()]
        nums = [self._assign(ref_id) for ref_id in ids]
        return "[" + ",".join(str(n) for n in nums) + "]"

    def replace(self, text: str) -> str:
        def repl(match: re.Match[str]) -> str:
            return self.cite_text(match.group(1))

        return re.sub(r"\{\{\s*cite:([^}]+)\}\}", repl, str(text))

    def ordered_references(self) -> list[dict[str, Any]]:
        used = [self.by_id[ref_id] for ref_id in self.order if ref_id in self.by_id]
        used_ids = {str(ref.get("id")) for ref in used}
        unused = [ref for ref in self.references if str(ref.get("id")) not in used_ids]
        return used + unused


class NumberingManager:
    def __init__(self, source: dict[str, Any]) -> None:
        self.figure_nums: dict[str, int] = {}
        self.table_nums: dict[str, int] = {}
        self.equation_nums: dict[str, int] = {}
        self._collect(source)

    @staticmethod
    def _block_id(block: Any) -> str | None:
        if isinstance(block, str):
            return None
        if not isinstance(block, dict):
            return None
        return block.get("id") or block.get("ref") or block.get("figure") or block.get("table") or block.get("equation")

    def _assign(self, bucket: dict[str, int], item_id: str | None) -> None:
        if item_id and item_id not in bucket:
            bucket[item_id] = len(bucket) + 1

    def _collect(self, source: dict[str, Any]) -> None:
        for section in source.get("sections", []) or []:
            content = section.get("content")
            if content:
                for block in content:
                    if not isinstance(block, dict):
                        continue
                    kind = str(block.get("type", "")).lower()
                    item_id = self._block_id(block)
                    if kind in {"figure", "fig"}:
                        self._assign(self.figure_nums, item_id)
                    elif kind in {"table", "tbl"}:
                        self._assign(self.table_nums, item_id)
                    elif kind in {"equation", "eq", "formula"}:
                        self._assign(self.equation_nums, item_id)
                continue
            for fig in section.get("figures", []) or []:
                self._assign(self.figure_nums, self._block_id(fig))
            for tbl in section.get("tables", []) or []:
                self._assign(self.table_nums, self._block_id(tbl))
            for eq in section.get("equations", []) or []:
                self._assign(self.equation_nums, self._block_id(eq))

    def replace(self, text: str) -> str:
        def repl(match: re.Match[str]) -> str:
            kind, item_id = match.group(1).lower(), match.group(2).strip()
            if kind in {"fig", "figure"}:
                return f"图{self.figure_nums.get(item_id, '?')}"
            if kind in {"tbl", "table"}:
                return f"表{self.table_nums.get(item_id, '?')}"
            return f"式({self.equation_nums.get(item_id, '?')})"

        return re.sub(r"\{\{\s*(fig|figure|tbl|table|eq|equation):([^}]+)\}\}", repl, str(text))


class HeadingNumberer:
    def __init__(self) -> None:
        self.counters = [0, 0, 0]

    def next(self, level: int) -> str:
        if level < 1 or level > 3:
            return ""
        self.counters[level - 1] += 1
        for i in range(level, 3):
            self.counters[i] = 0
        if level == 2 and self.counters[0] == 0:
            self.counters[0] = 1
        if level == 3 and self.counters[1] == 0:
            self.counters[1] = 1
        return ".".join(str(n) for n in self.counters[:level])


def normalized_text(text: str, citations: CitationManager, numbering: NumberingManager) -> str:
    return citations.replace(numbering.replace(str(text)))


def pre_register_appendix_citations(appendices: list[dict[str, Any]], citations: CitationManager, numbering: NumberingManager) -> None:
    for appendix in appendices or []:
        for text in appendix.get("paragraphs", []) or []:
            normalized_text(str(text), citations, numbering)


def text_parts_with_citations(text: str, citations: CitationManager, numbering: NumberingManager) -> list[tuple[str, bool]]:
    text = numbering.replace(str(text))
    parts: list[tuple[str, bool]] = []
    pos = 0
    for match in re.finditer(r"\{\{\s*cite:([^}]+)\}\}", text):
        if match.start() > pos:
            parts.append((text[pos : match.start()], False))
        parts.append((citations.cite_text(match.group(1)), True))
        pos = match.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    return [(part, superscript) for part, superscript in parts if part]


def maps_by_id(items: list[dict[str, Any]] | None) -> dict[str, dict[str, Any]]:
    return {str(item.get("id")): item for item in (items or []) if item.get("id") is not None}


def resolve_asset(path_text: str | None, source_path: Path) -> Path | None:
    if not path_text:
        return None
    path = Path(path_text)
    if path.is_absolute() and path.exists():
        return path
    candidates = [source_path.parent / path, Path.cwd() / path]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def placeholder_image_stream(label: str) -> BytesIO:
    stream = BytesIO()
    if Image is None or ImageDraw is None:
        raise RuntimeError("Pillow is required to generate a missing-figure placeholder")
    image = Image.new("RGB", (900, 520), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((16, 16, 884, 504), outline=(60, 60, 60), width=3)
    draw.line((16, 504, 884, 16), fill=(180, 180, 180), width=2)
    draw.line((16, 16, 884, 504), fill=(180, 180, 180), width=2)
    draw.text((38, 38), f"Missing figure: {label}", fill=(40, 40, 40))
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def set_inline_shape_size(shape, width_mm: float | None = None, max_width_mm: float | None = None) -> None:
    if width_mm is None:
        width_mm = max_width_mm
    if width_mm is None:
        return
    if max_width_mm is not None:
        width_mm = min(width_mm, max_width_mm)
    old_width = shape.width
    old_height = shape.height
    shape.width = int(width_mm * EMU_PER_MM)
    if old_width:
        shape.height = int(old_height * shape.width / old_width)


def add_figure(
    doc: Document,
    fig: dict[str, Any],
    *,
    number: int,
    source_path: Path,
    style: dict[str, Any],
    citations: CitationManager,
    numbering: NumberingManager,
) -> None:
    max_width = float(fig.get("max_width_mm") or style_at(style, "figures", "half_column_max_width_mm", default=75))
    width = fig.get("width_mm")
    image_path = resolve_asset(fig.get("path"), source_path)
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = paragraph.add_run()
    if image_path:
        shape = run.add_picture(str(image_path))
    else:
        print(f"WARNING: figure asset not found for {fig.get('id', number)!r}; inserted an inline placeholder image")
        shape = run.add_picture(placeholder_image_stream(str(fig.get("id", number))))
    set_inline_shape_size(shape, float(width) if width else None, max_width)

    caption = normalized_text(fig.get("caption", ""), citations, numbering)
    add_strict_paragraph(doc, f"图{number}  {caption}", style, "figure_caption")
    note = fig.get("note") or fig.get("notes")
    if note:
        add_strict_paragraph(
            doc,
            normalized_text(note if isinstance(note, str) else "；".join(map(str, note)), citations, numbering),
            style,
            "figure_note",
        )


def set_table_cell_margins(table, margin_twips: int = 108) -> None:
    tbl_pr = table._tbl.tblPr
    cell_mar = get_or_add(tbl_pr, "w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        node = get_or_add(cell_mar, f"w:{edge}")
        node.set(qn("w:w"), str(margin_twips if edge in {"left", "right"} else 0))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, style: dict[str, Any]) -> None:
    tbl_pr = table._tbl.tblPr
    width = style_at(style, "strict_layout", "tables", "tblW", default={}) or {}
    if not width:
        return
    tbl_w = get_or_add(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(width.get("w", "5536")))
    tbl_w.set(qn("w:type"), str(width.get("type", "dxa")))


def set_table_layout_fixed(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = get_or_add(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")


def set_cell_border(cell, **edges: dict[str, str]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = get_or_add(tc_pr, "w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        data = edges.get(edge)
        if data is None:
            remove_child(borders, f"w:{edge}")
            continue
        element = get_or_add(borders, f"w:{edge}")
        element.set(qn("w:val"), data.get("val", "single"))
        if data.get("val") not in {"nil", "none"}:
            element.set(qn("w:sz"), str(data.get("sz", 4)))
            element.set(qn("w:space"), str(data.get("space", 0)))
            element.set(qn("w:color"), data.get("color", "000000"))


def apply_three_line_borders(table) -> None:
    row_count = len(table.rows)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            edges: dict[str, dict[str, str]] = {}
            if row_idx == 0:
                edges["top"] = {"val": "single", "sz": 8}
                edges["bottom"] = {"val": "single", "sz": 4}
            if row_idx == row_count - 1:
                edges["bottom"] = {"val": "single", "sz": 8}
            set_cell_border(cell, **edges)


def set_cell_text(
    cell,
    text: Any,
    *,
    size_pt: float,
    sz_cs_pt: float | None = None,
    zh: str = DEFAULT_ZH_FONT,
    latin: str = DEFAULT_LATIN_FONT,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.CENTER,
    bold: bool | None = None,
) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    set_paragraph_format(paragraph, align=align)
    add_runs_from_text(paragraph, str(text), size_pt=size_pt, sz_cs_pt=sz_cs_pt, zh=zh, latin=latin, bold=bold)


def add_data_table(
    doc: Document,
    table_data: dict[str, Any],
    *,
    number: int,
    style: dict[str, Any],
    citations: CitationManager,
    numbering: NumberingManager,
) -> None:
    caption = normalized_text(table_data.get("caption", ""), citations, numbering)
    add_strict_paragraph(doc, f"表{number}  {caption}", style, "table_caption", keep_with_next=True)

    headers = [normalized_text(x, citations, numbering) for x in table_data.get("headers", [])]
    rows = table_data.get("rows", []) or []
    if not headers and rows:
        headers = ["" for _ in range(max(len(row) for row in rows))]
    col_count = max([len(headers)] + [len(row) for row in rows] + [1])
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, style)
    set_table_layout_fixed(table)
    set_table_cell_margins(table)
    apply_three_line_borders(table)

    body_size = strict_size_pt(style, "table_body", 7.5)
    table_body = strict_style(style, "table_body")
    body_sz_cs = float(table_body.get("szCs", table_body.get("sz", 15))) / 2
    for col_idx in range(col_count):
        value = headers[col_idx] if col_idx < len(headers) else ""
        set_cell_text(table.cell(0, col_idx), value, size_pt=body_size, sz_cs_pt=body_sz_cs)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx in range(col_count):
            value = row[col_idx] if col_idx < len(row) else ""
            set_cell_text(table.cell(row_idx, col_idx), normalized_text(value, citations, numbering), size_pt=body_size, sz_cs_pt=body_sz_cs)

    note = table_data.get("note") or table_data.get("notes")
    if note:
        note_text = note if isinstance(note, str) else " ".join(str(x) for x in note)
        add_strict_paragraph(doc, normalized_text(note_text, citations, numbering), style, "figure_note")


def add_equation(
    doc: Document,
    eq: dict[str, Any],
    *,
    number: int,
    style: dict[str, Any],
    citations: CitationManager,
    numbering: NumberingManager,
) -> None:
    eq_size = first_number(style, "equation", "font_size_pt", "font_size_pt_observed", default=9)
    text = normalized_text(eq.get("text") or eq.get("latex") or eq.get("omml") or "", citations, numbering)
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER)
    if eq.get("omml"):
        paragraph._p.append(parse_xml(str(eq["omml"])))
    else:
        omml_text = html.escape(text)
        paragraph._p.append(
            parse_xml(
                '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                f"<m:oMath><m:r><m:t>{omml_text}</m:t></m:r></m:oMath>"
                "</m:oMathPara>"
            )
        )
    # The template stores the displayed equation and the right-side number as
    # separate paragraphs. This avoids placing equations in tables/text boxes.
    add_strict_paragraph(doc, f"({number})", style, "equation_number")


def add_symbols(doc: Document, symbols: list[dict[str, Any]], style: dict[str, Any]) -> None:
    if not symbols:
        return
    add_strict_paragraph(doc, "符号说明", style, "table_caption")
    table = doc.add_table(rows=len(symbols), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, style)
    set_table_layout_fixed(table)
    set_table_cell_margins(table)
    for row_idx, item in enumerate(symbols):
        set_cell_text(table.cell(row_idx, 0), item.get("symbol", ""), size_pt=9, align=WD_ALIGN_PARAGRAPH.RIGHT)
        set_cell_text(table.cell(row_idx, 1), "——", size_pt=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        meaning = item.get("meaning", "")
        unit = item.get("unit", "")
        text = f"{meaning}，{unit}" if unit else str(meaning)
        set_cell_text(table.cell(row_idx, 2), text, size_pt=9, align=WD_ALIGN_PARAGRAPH.LEFT)
        for cell in table.rows[row_idx].cells:
            set_cell_border(cell)


def render_heading(doc: Document, section: dict[str, Any], heading_numbers: HeadingNumberer, style: dict[str, Any]) -> None:
    level = int(section.get("level", 0) or 0)
    title = str(section.get("title", "") or "")
    if level <= 0 or not title:
        return
    number = heading_numbers.next(level)
    if level == 1:
        add_strict_paragraph(doc, f"{number}  {title}", style, "heading_1", keep_with_next=True)
    elif level == 2:
        add_strict_paragraph(doc, f"{number}  {title}", style, "heading_2", bold=True, keep_with_next=True)
    elif level == 3:
        add_strict_paragraph(doc, f"{number}  {title}", style, "heading_3", keep_with_next=True)


def render_body_paragraph(
    doc: Document,
    text: str,
    style: dict[str, Any],
    citations: CitationManager | None = None,
    numbering: NumberingManager | None = None,
) -> None:
    body_style = strict_style(style, "body")
    size = strict_size_pt(style, "body", 10.5)
    sz_cs_pt = float(body_style.get("szCs", body_style.get("sz", 21))) / 2
    paragraph = doc.add_paragraph()
    set_raw_paragraph_attrs(paragraph, style, "body")
    if citations is not None and numbering is not None:
        for part, superscript in text_parts_with_citations(text, citations, numbering):
            run = paragraph.add_run(part)
            set_run_font(run, size_pt=size, sz_cs_pt=sz_cs_pt, superscript=superscript)
    else:
        add_runs_from_text(paragraph, text, size_pt=size, sz_cs_pt=sz_cs_pt, zh=DEFAULT_ZH_FONT, latin=DEFAULT_LATIN_FONT)


def render_block(
    doc: Document,
    block: Any,
    *,
    source_path: Path,
    style: dict[str, Any],
    figures: dict[str, dict[str, Any]],
    tables: dict[str, dict[str, Any]],
    equations: dict[str, dict[str, Any]],
    citations: CitationManager,
    numbering: NumberingManager,
) -> None:
    if isinstance(block, str):
        render_body_paragraph(doc, block, style, citations, numbering)
        return
    if not isinstance(block, dict):
        return
    kind = str(block.get("type", "paragraph")).lower()
    if kind in {"paragraph", "text", "p"}:
        render_body_paragraph(doc, block.get("text", ""), style, citations, numbering)
    elif kind in {"figure", "fig"}:
        item_id = block.get("id") or block.get("ref") or block.get("figure")
        fig = {**figures.get(str(item_id), {}), **block}
        add_figure(
            doc,
            fig,
            number=numbering.figure_nums.get(str(item_id), 0),
            source_path=source_path,
            style=style,
            citations=citations,
            numbering=numbering,
        )
    elif kind in {"table", "tbl"}:
        item_id = block.get("id") or block.get("ref") or block.get("table")
        table_data = {**tables.get(str(item_id), {}), **block}
        add_data_table(
            doc,
            table_data,
            number=numbering.table_nums.get(str(item_id), 0),
            style=style,
            citations=citations,
            numbering=numbering,
        )
    elif kind in {"equation", "eq", "formula"}:
        item_id = block.get("id") or block.get("ref") or block.get("equation")
        eq = {**equations.get(str(item_id), {}), **block}
        add_equation(
            doc,
            eq,
            number=numbering.equation_nums.get(str(item_id), 0),
            style=style,
            citations=citations,
            numbering=numbering,
        )
    elif kind in {"conclusion", "item"}:
        render_body_paragraph(doc, block.get("text", ""), style, citations, numbering)


def render_sections(
    doc: Document,
    source: dict[str, Any],
    *,
    source_path: Path,
    style: dict[str, Any],
    citations: CitationManager,
    numbering: NumberingManager,
) -> None:
    figures = maps_by_id(source.get("figures"))
    tables = maps_by_id(source.get("tables"))
    equations = maps_by_id(source.get("equations"))
    heading_numbers = HeadingNumberer()

    for section in source.get("sections", []) or []:
        render_heading(doc, section, heading_numbers, style)
        content = section.get("content")
        if content:
            for block in content:
                render_block(
                    doc,
                    block,
                    source_path=source_path,
                    style=style,
                    figures=figures,
                    tables=tables,
                    equations=equations,
                    citations=citations,
                    numbering=numbering,
                )
            continue

        for text in section.get("paragraphs", []) or []:
            render_body_paragraph(doc, text, style, citations, numbering)
        for eq_ref in section.get("equations", []) or []:
            render_block(
                doc,
                {"type": "equation", **(eq_ref if isinstance(eq_ref, dict) else {"id": eq_ref})},
                source_path=source_path,
                style=style,
                figures=figures,
                tables=tables,
                equations=equations,
                citations=citations,
                numbering=numbering,
            )
        for fig_ref in section.get("figures", []) or []:
            render_block(
                doc,
                {"type": "figure", **(fig_ref if isinstance(fig_ref, dict) else {"id": fig_ref})},
                source_path=source_path,
                style=style,
                figures=figures,
                tables=tables,
                equations=equations,
                citations=citations,
                numbering=numbering,
            )
        for tbl_ref in section.get("tables", []) or []:
            render_block(
                doc,
                {"type": "table", **(tbl_ref if isinstance(tbl_ref, dict) else {"id": tbl_ref})},
                source_path=source_path,
                style=style,
                figures=figures,
                tables=tables,
                equations=equations,
                citations=citations,
                numbering=numbering,
            )
        for idx, conclusion in enumerate(section.get("conclusions", []) or [], start=1):
            render_body_paragraph(doc, f"（{idx}）{conclusion}", style, citations, numbering)


def render_appendices(doc: Document, appendices: list[dict[str, Any]], style: dict[str, Any], citations: CitationManager, numbering: NumberingManager) -> None:
    for idx, appendix in enumerate(appendices or [], start=1):
        letter = appendix.get("letter") or chr(ord("A") + idx - 1)
        title = appendix.get("title", "")
        heading = f"附录{letter}" + (f"  {title}" if title else "")
        add_strict_paragraph(doc, heading, style, "appendix", keep_with_next=True)
        for text in appendix.get("paragraphs", []) or []:
            render_body_paragraph(doc, text, style, citations, numbering)


def has_cjk(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text))


def format_author(author: Any, *, english: bool | None = None) -> str:
    if isinstance(author, str):
        return author.strip()
    if not isinstance(author, dict):
        return str(author).strip()
    family = str(author.get("family") or author.get("surname") or "").strip()
    given = str(author.get("given") or author.get("name") or "").strip()
    literal = str(author.get("literal") or "").strip()
    if literal:
        return literal
    if english is None:
        english = not has_cjk(family + given)
    if english:
        return (family.upper() + (" " + given if given else "")).strip()
    return (family + given).strip()


def format_authors(authors: Any) -> str:
    if isinstance(authors, str):
        return authors.strip()
    authors = list(authors or [])
    if not authors:
        return ""
    rendered_all = [format_author(author) for author in authors]
    english = not has_cjk("".join(rendered_all))
    if len(rendered_all) > 3:
        suffix = "et al" if english else "等"
        rendered_all = rendered_all[:3] + [suffix]
    return ", ".join(rendered_all)


def ensure_period(text: str) -> str:
    text = re.sub(r"\s+", " ", str(text).strip())
    return text if text.endswith(".") else text + "."


def fmt_pages(ref: dict[str, Any]) -> str:
    pages = ref.get("pages") or ref.get("page")
    return f": {pages}" if pages else ""


def format_reference_text(ref: dict[str, Any]) -> str:
    if ref.get("text"):
        return ensure_period(str(ref["text"]))
    ref_type = str(ref.get("type") or ref.get("entry_type") or "J").upper()
    authors = format_authors(ref.get("authors") or ref.get("applicants") or ref.get("organization"))
    title = str(ref.get("title") or "").strip()
    year = ref.get("year")
    if ref_type == "J":
        journal = str(ref.get("journal") or ref.get("source") or "").strip()
        volume = str(ref.get("volume") or "").strip()
        issue = str(ref.get("issue") or "").strip()
        vol_issue = volume + (f"({issue})" if issue else "")
        tail = f"{year}, {vol_issue}{fmt_pages(ref)}" if year and vol_issue else str(year or "")
        return ensure_period(f"{authors}. {title}[J]. {journal}, {tail}")
    if ref_type == "M":
        place = str(ref.get("place") or "").strip()
        publisher = str(ref.get("publisher") or "").strip()
        edition = str(ref.get("edition") or "").strip()
        edition_part = f"{edition}. " if edition else ""
        return ensure_period(f"{authors}. {title}[M]. {edition_part}{place}: {publisher}, {year}{fmt_pages(ref)}")
    if ref_type == "C":
        booktitle = str(ref.get("booktitle") or ref.get("conference") or "").strip()
        place = str(ref.get("place") or "").strip()
        publisher = str(ref.get("publisher") or "").strip()
        pub = f"{place}: {publisher}, {year}{fmt_pages(ref)}" if publisher else f"{place}, {year}{fmt_pages(ref)}"
        return ensure_period(f"{authors}. {title}[C]//{booktitle}. {pub}")
    if ref_type == "P":
        number = str(ref.get("number") or ref.get("patent_number") or "").strip()
        date = str(ref.get("date") or ref.get("published") or "").strip()
        return ensure_period(f"{authors}. {title}: {number}[P]. {date}")
    if ref_type == "S":
        number = str(ref.get("number") or ref.get("standard_number") or "").strip()
        place = str(ref.get("place") or "").strip()
        publisher = str(ref.get("publisher") or "").strip()
        return ensure_period(f"{authors}. {title}: {number}[S]. {place}: {publisher}, {year}")
    if ref_type == "D":
        place = str(ref.get("place") or "").strip()
        publisher = str(ref.get("publisher") or ref.get("institution") or "").strip()
        return ensure_period(f"{authors}. {title}[D]. {place}: {publisher}, {year}")
    if ref_type == "R":
        place = str(ref.get("place") or "").strip()
        publisher = str(ref.get("publisher") or "").strip()
        return ensure_period(f"{authors}. {title}[R]. {place}: {publisher}, {year}{fmt_pages(ref)}")
    if ref_type == "N":
        newspaper = str(ref.get("newspaper") or ref.get("source") or "").strip()
        date = str(ref.get("date") or "").strip()
        page = str(ref.get("page") or "").strip()
        page_part = f"({page})" if page else ""
        return ensure_period(f"{authors}. {title}[N]. {newspaper}, {date}{page_part}")
    if ref_type in {"EB", "EB/OL"}:
        date = str(ref.get("date") or "").strip()
        cited = str(ref.get("cited") or ref.get("accessed") or "").strip()
        url = str(ref.get("url") or "").strip()
        date_part = f"({date}) " if date else ""
        cited_part = f"[{cited}]. " if cited else ""
        return ensure_period(f"{authors}. {title}[EB/OL]. {date_part}{cited_part}{url}")
    return ensure_period(f"{authors}. {title}[{ref_type}].")


def reference_paragraphs(ref: dict[str, Any]) -> list[str]:
    lines = [format_reference_text(ref)]
    for key in ("text_en", "translated_text", "english_text"):
        if ref.get(key):
            lines.append(ensure_period(str(ref[key])))
            break
    return lines


def build_doc(source: dict[str, Any], style: dict[str, Any], source_path: Path) -> Document:
    doc = Document()
    set_document_defaults(doc, style)
    apply_page_setup(doc, style)

    citations = CitationManager(source.get("references", []) or [])
    numbering = NumberingManager(source)

    add_strict_paragraph(doc, source.get("article_type", "研究开发"), style, "article_type")

    doi = source.get("doi", "")
    if doi:
        add_strict_labeled_paragraph(doc, "DOI", f"：{doi}", style, "doi", label_bold=True)

    add_strict_paragraph(doc, source.get("title_zh", ""), style, "zh_title")
    authors = source.get("authors", []) or []
    build_author_paragraph(doc, authors, english=False, style=style)
    aff_text = affiliation_line(source.get("affiliations", []) or [], english=False)
    if aff_text:
        add_paragraph(
            doc,
            aff_text,
            size_pt=strict_size_pt(style, "zh_affiliations", 9),
            sz_cs_pt=float(strict_style(style, "zh_affiliations").get("szCs", 18)) / 2,
            zh=strict_font_zh(style, "zh_affiliations"),
            latin=strict_font_latin(style, "zh_affiliations"),
        )
        set_raw_paragraph_attrs(doc.paragraphs[-1], style, "zh_affiliations")

    for note in source.get("front_matter_notes", []) or []:
        add_strict_paragraph(doc, str(note), style, "front_matter_note")

    add_strict_labeled_paragraph(
        doc,
        "摘要：",
        normalized_text(source.get("abstract_zh", ""), citations, numbering),
        style,
        "zh_abstract",
    )
    add_strict_labeled_paragraph(
        doc,
        "关键词",
        "：" + "；".join(source.get("keywords_zh", []) or []),
        style,
        "zh_keywords",
    )
    classification = source.get("classification", {}) or {}
    add_strict_classification_line(doc, classification, style)

    add_strict_paragraph(doc, source.get("title_en", ""), style, "en_title", bold=True)
    build_author_paragraph(doc, authors, english=True, style=style)
    en_aff_text = affiliation_line(source.get("affiliations", []) or [], english=True)
    if en_aff_text:
        add_paragraph(
            doc,
            en_aff_text,
            size_pt=strict_size_pt(style, "en_affiliations", 9),
            sz_cs_pt=float(strict_style(style, "en_affiliations").get("szCs", 18)) / 2,
            zh=strict_font_zh(style, "en_affiliations"),
            latin=strict_font_latin(style, "en_affiliations"),
        )
        set_raw_paragraph_attrs(doc.paragraphs[-1], style, "en_affiliations")
    add_strict_labeled_paragraph(
        doc,
        "Abstract:",
        " " + normalized_text(source.get("abstract_en", ""), citations, numbering),
        style,
        "en_abstract",
        label_bold=True,
    )
    add_strict_labeled_paragraph(
        doc,
        "Keywords:",
        " " + "; ".join(source.get("keywords_en", []) or []),
        style,
        "en_keywords",
        label_bold=True,
    )

    render_sections(doc, source, source_path=source_path, style=style, citations=citations, numbering=numbering)
    add_symbols(doc, source.get("symbols", []) or [], style)

    add_strict_paragraph(doc, "参考文献", style, "table_caption")
    pre_register_appendix_citations(source.get("appendices", []) or [], citations, numbering)
    for idx, ref in enumerate(citations.ordered_references(), start=1):
        ref_lines = reference_paragraphs(ref)
        add_strict_paragraph(doc, f"[{idx}]　{ref_lines[0]}", style, "references")
        for extra in ref_lines[1:]:
            add_strict_paragraph(doc, extra, style, "references")

    render_appendices(doc, source.get("appendices", []) or [], style, citations, numbering)
    return doc


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--style", type=Path, default=Path("STYLE_MAP.json"))
    parser.add_argument("--out", type=Path, default=Path("output/manuscript.docx"))
    args = parser.parse_args()

    source = load_structured(args.source)
    style = load_style(args.style)
    doc = build_doc(source, style, args.source.resolve())
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
