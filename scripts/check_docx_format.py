#!/usr/bin/env python3
"""Check a generated manuscript DOCX against the local journal style map.

The checker is intentionally conservative. It reports machine-checkable
violations as ERROR and uncertain visual/content checks as WARN with concrete
locations and suggested fixes.
"""
from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}

STATUS_RANK = {"PASS": 0, "INFO": 1, "WARN": 2, "ERROR": 3}
MANDATORY_CATEGORIES = [
    "页面设置",
    "禁用 Word 特性",
    "标题编号",
    "图题位置",
    "表题位置",
    "编号连续性",
    "正文呼应",
    "参考文献",
    "域代码/超链接",
    "三线表",
]
EXTRA_CATEGORIES = ["字体字号", "细排版", "前置换行结构", "表格细节", "图片嵌入", "公式", "附录", "注释"]
CONTENT_CATEGORIES = ["标题摘要关键词", "语言风格", "单位符号", "结论结语"]


def cm(value) -> float | None:
    return value.cm if value is not None else None


def near(actual: float | None, expected: float, tol: float = 0.08) -> bool:
    return actual is not None and abs(actual - expected) <= tol


def load_style(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    data["_style_base_dir"] = str(path.resolve().parent)
    return data


def style_at(style: dict[str, Any], *path: str, default: Any = None) -> Any:
    cur: Any = style
    for key in path:
        if not isinstance(cur, dict) or key not in cur:
            return default
        cur = cur[key]
    return cur


def resolve_style_path(style: dict[str, Any], value: str | None) -> Path | None:
    if not value:
        return None
    path = Path(value)
    if path.is_absolute():
        return path
    base = Path(str(style.get("_style_base_dir", ".")))
    return base / path


def load_standard_keywords(style: dict[str, Any]) -> tuple[set[str], str, str]:
    rules = style_at(style, "content_rules", "keywords_zh", default={}) or {}
    source = rules.get("standard_keyword_file") or rules.get("standard_keyword_source")
    path = resolve_style_path(style, source)
    if path is None:
        return set(), "", "not configured"
    if not path.exists():
        return set(), str(path), "missing"

    try:
        if path.suffix.lower() == ".json":
            data = json.loads(path.read_text(encoding="utf-8"))
            words = data.get("keywords", [])
        else:
            words = [line.strip() for line in path.read_text(encoding="utf-8").splitlines()]
    except Exception as exc:
        return set(), str(path), f"load failed: {exc}"

    keywords = {str(word).strip() for word in words if str(word).strip()}
    if not keywords:
        return set(), str(path), "empty"
    return keywords, str(path), "loaded"


def add(items: list[dict[str, str]], category: str, status: str, message: str, location: str = "", fix: str = "") -> None:
    items.append(
        {
            "category": category,
            "status": status,
            "message": message,
            "location": location,
            "fix": fix,
        }
    )


def summarize_category(items: list[dict[str, str]], category: str) -> str:
    status = "PASS"
    found = False
    for item in items:
        if item["category"] != category:
            continue
        found = True
        if STATUS_RANK[item["status"]] > STATUS_RANK[status]:
            status = item["status"]
    return status if found else "INFO"


def read_docx_parts(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        parts: dict[str, str] = {}
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    parts[name] = package.read(name).decode("utf-8", errors="ignore")
                except Exception:
                    parts[name] = ""
    return {
        "names": names,
        "parts": parts,
        "document_xml": parts.get("word/document.xml", ""),
    }


def xml_root(parts: dict[str, Any]) -> etree._Element:
    return etree.fromstring(parts["document_xml"].encode("utf-8"))


def local_attrs(element) -> dict[str, str]:
    if element is None:
        return {}
    return {str(key).split("}")[-1]: value for key, value in element.attrib.items()}


def xml_text(element) -> str:
    return "".join(t.text or "" for t in element.xpath(".//w:t", namespaces=NS)).strip()


def body_paragraph_elements(parts: dict[str, Any]) -> list[etree._Element]:
    return xml_root(parts).xpath("//w:body/w:p", namespaces=NS)


def first_text_run(paragraph_xml: etree._Element) -> etree._Element | None:
    for run in paragraph_xml.xpath("./w:r", namespaces=NS):
        if xml_text(run):
            return run
    return None


def text_runs(paragraph_xml: etree._Element) -> list[etree._Element]:
    return [run for run in paragraph_xml.xpath("./w:r", namespaces=NS) if xml_text(run)]


def ppr_child_attrs(paragraph_xml: etree._Element, child: str) -> dict[str, str]:
    p_pr = paragraph_xml.find(qn("w:pPr"))
    if p_pr is None:
        return {}
    return local_attrs(p_pr.find(qn(f"w:{child}")))


def rpr_child_attrs(run_xml: etree._Element | None, child: str) -> dict[str, str]:
    if run_xml is None:
        return {}
    r_pr = run_xml.find(qn("w:rPr"))
    if r_pr is None:
        return {}
    return local_attrs(r_pr.find(qn(f"w:{child}")))


def rpr_has_child(run_xml: etree._Element | None, child: str) -> bool:
    if run_xml is None:
        return False
    r_pr = run_xml.find(qn("w:rPr"))
    return r_pr is not None and r_pr.find(qn(f"w:{child}")) is not None


def font_token(value: str | None) -> str:
    return str(value or "").split(";")[0].strip()


def raw_int(value: str | None) -> int | None:
    try:
        return int(str(value))
    except Exception:
        return None


def near_twips(actual: str | None, expected: str | int | None, tol: int) -> bool:
    if expected is None:
        return True
    left = raw_int(actual)
    right = raw_int(str(expected))
    return left is not None and right is not None and abs(left - right) <= tol


def real_note_ids(xml: str, tag: str) -> list[int]:
    ids: list[int] = []
    for match in re.finditer(rf"<w:{tag}\b[^>]*w:id=\"(-?\d+)\"", xml):
        note_id = int(match.group(1))
        if note_id > 0:
            ids.append(note_id)
    return ids


def check_page(doc: Document, style: dict[str, Any], parts: dict[str, Any], items: list[dict[str, str]]) -> None:
    before = len(items)
    section = doc.sections[0]
    expected = style.get("page_setup", {})
    checks = [
        ("纸张宽度", cm(section.page_width), expected.get("width_cm", 21.0), "设置为 A4 宽 21.0 cm"),
        ("纸张高度", cm(section.page_height), expected.get("height_cm", 29.7), "设置为 A4 高 29.7 cm"),
        ("上边距", cm(section.top_margin), expected.get("top_margin_cm", 2.5), "设置上边距 2.5 cm"),
        ("下边距", cm(section.bottom_margin), expected.get("bottom_margin_cm", 2.0), "设置下边距 2.0 cm"),
        ("左边距", cm(section.left_margin), expected.get("left_margin_cm", 2.0), "设置左边距 2.0 cm"),
        ("右边距", cm(section.right_margin), expected.get("right_margin_cm", 2.0), "设置右边距 2.0 cm"),
    ]
    for label, actual, target, fix in checks:
        if not near(actual, float(target)):
            add(items, "页面设置", "ERROR", f"{label}不匹配：实际 {actual:.2f} cm，期望 {target:.2f} cm", fix=fix)

    xml = parts["document_xml"]
    if re.search(r"<w:cols\b[^>]*w:num=\"[2-9]", xml):
        add(items, "页面设置", "ERROR", "检测到多栏设置，模板要求单栏", fix="删除分栏，使用单栏正文")
    grid = re.search(r"<w:docGrid\b([^>]*)>", xml)
    if not grid:
        add(items, "页面设置", "WARN", "未检测到 docGrid，可能没有落实 44 行/46 字页面网格", fix="按模板设置 linesAndChars 文档网格")
    else:
        attrs = grid.group(1)
        if 'w:type="linesAndChars"' not in attrs:
            add(items, "页面设置", "WARN", "docGrid 不是 linesAndChars", fix="设置每页 44 行、每行 46 字")
        if 'w:linePitch="324"' not in attrs:
            add(items, "页面设置", "WARN", "docGrid linePitch 与模板观测值不一致", fix="使用模板观测 linePitch=324")

    strict_page = style_at(style, "strict_layout", "page_ooxml", default={}) or {}
    if strict_page:
        sect_prs = xml_root(parts).xpath("//w:sectPr", namespaces=NS)
        sect_pr = sect_prs[-1] if sect_prs else None
        tol = int(style_at(style, "strict_layout", "tolerance", "twips", default=8))
        for child_name, expected_attrs in strict_page.items():
            actual_attrs = local_attrs(sect_pr.find(qn(f"w:{child_name}"))) if sect_pr is not None else {}
            for key, expected_value in (expected_attrs or {}).items():
                actual_value = actual_attrs.get(key)
                if key in {"w", "h", "top", "bottom", "left", "right", "header", "footer", "gutter", "linePitch"}:
                    ok = near_twips(actual_value, expected_value, tol)
                else:
                    ok = str(actual_value) == str(expected_value)
                if not ok:
                    add(
                        items,
                        "页面设置",
                        "ERROR",
                        f"{child_name}/{key} raw OOXML 不匹配：实际 {actual_value}，期望 {expected_value}",
                        fix="按 strict_layout.page_ooxml 写入模板原始页面属性",
                    )

    if parts["document_xml"].count("<w:sectPr") > 1:
        add(items, "页面设置", "WARN", "检测到多个 sectPr，模板通常不需要分节符", fix="检查并删除误插入分节符")

    if len(items) == before:
        add(items, "页面设置", "PASS", "页面尺寸、边距、单栏和页面网格未发现机器可检问题")


def has_real_header_footer(parts: dict[str, Any], prefix: str) -> bool:
    for name, xml in parts["parts"].items():
        if name.startswith(f"word/{prefix}") and re.search(r"<w:t[^>]*>\s*[^<\s]", xml):
            return True
    return False


def check_forbidden_features(parts: dict[str, Any], items: list[dict[str, str]]) -> None:
    names = parts["names"]
    document_xml = parts["document_xml"]
    footnote_ids = real_note_ids(parts["parts"].get("word/footnotes.xml", ""), "footnote")
    endnote_ids = real_note_ids(parts["parts"].get("word/endnotes.xml", ""), "endnote")

    forbidden_problem = False
    forbidden_checks = [
        ("文本框", "w:txbxContent" in document_xml or "v:textbox" in document_xml, "删除文本框，改用普通段落/内嵌图片"),
        ("批注", "word/comments.xml" in names, "接受/删除全部批注"),
        ("脚注", bool(footnote_ids), "提交稿不要保留脚注；首页说明改为普通文本"),
        ("尾注", bool(endnote_ids), "删除尾注，参考文献用纯文本编号"),
        ("页眉", has_real_header_footer(parts, "header"), "删除页眉"),
        ("页脚", has_real_header_footer(parts, "footer"), "删除页脚"),
        ("浮动图片", "<wp:anchor" in document_xml, "图片必须按嵌入型插入，不使用文字环绕"),
    ]
    for label, bad, fix in forbidden_checks:
        if bad:
            forbidden_problem = True
            add(items, "禁用 Word 特性", "ERROR", f"发现禁用或高风险 Word 特性：{label}", fix=fix)

    rels_with_hyperlink = any(
        name.endswith(".rels") and "relationships/hyperlink" in xml
        for name, xml in parts["parts"].items()
    )
    has_hyperlink = "<w:hyperlink" in document_xml or rels_with_hyperlink
    has_field = any(token in document_xml for token in ("w:fldChar", "w:instrText", "<w:fldSimple"))
    if has_hyperlink:
        add(items, "域代码/超链接", "ERROR", "检测到超链接", fix="将 DOI、网址和文献引用转为纯文本")
    if has_field:
        add(items, "域代码/超链接", "ERROR", "检测到 Word 域代码", fix="提交稿前平铺 EndNote/Zotero/REF/SEQ 等域代码")

    if not forbidden_problem:
        add(items, "禁用 Word 特性", "PASS", "未发现文本框、批注、脚注、尾注、页眉页脚或浮动图片")
    if not has_hyperlink and not has_field:
        add(items, "域代码/超链接", "PASS", "未发现域代码或超链接")


def paragraph_texts(doc: Document) -> list[str]:
    return [p.text.strip() for p in doc.paragraphs]


def heading_level(text: str) -> tuple[int, tuple[int, ...]] | None:
    if re.match(r"^\d+\.\d+\.\d+\s+\S", text):
        nums = tuple(int(x) for x in text.split()[0].split("."))
        return 3, nums
    if re.match(r"^\d+\.\d+\s+\S", text):
        nums = tuple(int(x) for x in text.split()[0].split("."))
        return 2, nums
    if re.match(r"^\d+\s+\S", text):
        return 1, (int(text.split()[0]),)
    return None


def check_headings(doc: Document, items: list[dict[str, str]]) -> None:
    before = len(items)
    expected = [0, 0, 0]
    found_h1 = False
    for idx, text in enumerate(paragraph_texts(doc)):
        if not text:
            continue
        if re.match(r"^[一二三四五六七八九十]+[、.．]", text):
            add(items, "标题编号", "ERROR", "发现中文数字标题编号", f"paragraph {idx}: {text[:80]}", "改为 1 / 1.1 / 1.1.1")
        parsed = heading_level(text)
        if not parsed:
            continue
        level, nums = parsed
        if level == 1:
            found_h1 = True
            expected[0] += 1
            expected[1] = 0
            expected[2] = 0
        elif level == 2:
            if expected[0] == 0:
                expected[0] = nums[0]
            expected[1] += 1
            expected[2] = 0
        else:
            if expected[0] == 0:
                expected[0] = nums[0]
            if expected[1] == 0:
                expected[1] = nums[1]
            expected[2] += 1
        target = tuple(expected[:level])
        if nums != target:
            add(items, "标题编号", "ERROR", f"标题编号顺序异常：检测到 {'.'.join(map(str, nums))}，期望 {'.'.join(map(str, target))}", f"paragraph {idx}: {text[:80]}", "让生成器自动编号或按层级重排")
        if not re.match(r"^\d+(?:\.\d+)*\s{2,}\S", text):
            add(items, "标题编号", "WARN", "标题编号与标题之间不是两个空格", f"paragraph {idx}: {text[:80]}", "格式应为：1  材料和方法")
    if not found_h1:
        add(items, "标题编号", "WARN", "未检测到一级标题", fix="正文一级标题使用 1  / 2  / 3  格式")
    if len(items) == before:
        add(items, "标题编号", "PASS", "一级/二级/三级标题编号未发现机器可检问题")


def iter_block_items(doc: Document):
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield "p", Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield "tbl", Table(child, doc)


def paragraph_has_image(paragraph) -> bool:
    xml = paragraph._p.xml
    return "<w:drawing" in xml or "<w:pict" in xml


def paragraph_has_math(paragraph) -> bool:
    return "<m:oMath" in paragraph._p.xml or "<m:oMathPara" in paragraph._p.xml


def caption_kind(text: str) -> tuple[str, int] | None:
    stripped = text.strip()
    match = re.match(r"^图\s*(\d+)\s+", stripped)
    if match:
        return "figure", int(match.group(1))
    match = re.match(r"^表\s*(\d+)\s+", stripped)
    if match:
        return "table", int(match.group(1))
    return None


def previous_nonempty(blocks: list[tuple[str, Any]], index: int) -> tuple[int, str, Any] | None:
    j = index - 1
    while j >= 0:
        kind, obj = blocks[j]
        if kind == "tbl":
            return j, kind, obj
        if obj.text.strip() or paragraph_has_image(obj) or paragraph_has_math(obj):
            return j, kind, obj
        j -= 1
    return None


def next_nonempty(blocks: list[tuple[str, Any]], index: int) -> tuple[int, str, Any] | None:
    j = index + 1
    while j < len(blocks):
        kind, obj = blocks[j]
        if kind == "tbl":
            return j, kind, obj
        if obj.text.strip() or paragraph_has_image(obj) or paragraph_has_math(obj):
            return j, kind, obj
        j += 1
    return None


def check_caption_positions(doc: Document, items: list[dict[str, str]]) -> list[tuple[int, Any, str]]:
    blocks = list(iter_block_items(doc))
    data_tables: list[tuple[int, Any, str]] = []
    seen_figure = False
    seen_table = False
    fig_issue = False
    table_issue = False

    for index, (kind, obj) in enumerate(blocks):
        if kind == "p" and paragraph_has_image(obj):
            seen_figure = True
            nxt = next_nonempty(blocks, index)
            if not nxt or nxt[1] != "p" or not re.match(r"^图\s*\d+\s+", nxt[2].text.strip()):
                fig_issue = True
                add(items, "图题位置", "ERROR", "图片后未紧跟图题，或图题格式不合格", f"block {index}", "图题必须在图下方，格式：图1  图题")
        if kind == "p" and re.match(r"^图\s*\d+\s+", obj.text.strip()):
            seen_figure = True
            prev = previous_nonempty(blocks, index)
            if not prev or prev[1] != "p" or not paragraph_has_image(prev[2]):
                fig_issue = True
                add(items, "图题位置", "WARN", "检测到图题，但其上一块不是图片", f"block {index}: {obj.text[:80]}", "图题应紧跟在对应图片下方")
            if not re.match(r"^图\s*\d+\s{2,}\S", obj.text.strip()):
                fig_issue = True
                add(items, "图题位置", "WARN", "图题编号后建议使用两个空格", f"block {index}: {obj.text[:80]}", "格式：图1  图题")

        if kind == "p" and re.match(r"^表\s*\d+\s+", obj.text.strip()):
            seen_table = True
            nxt = next_nonempty(blocks, index)
            if not nxt or nxt[1] != "tbl":
                table_issue = True
                add(items, "表题位置", "ERROR", "表题后未紧跟表格", f"block {index}: {obj.text[:80]}", "表题必须在表上方并紧邻表格")
            if not re.match(r"^表\s*\d+\s{2,}\S", obj.text.strip()):
                table_issue = True
                add(items, "表题位置", "WARN", "表题编号后建议使用两个空格", f"block {index}: {obj.text[:80]}", "格式：表1  表题")

        if kind == "tbl":
            seen_table = True
            prev = previous_nonempty(blocks, index)
            prev_text = prev[2].text.strip() if prev and prev[1] == "p" else ""
            if re.match(r"^表\s*\d+\s+", prev_text):
                data_tables.append((index, obj, prev_text))
            elif prev_text == "符号说明":
                add(items, "表题位置", "PASS", "符号说明表按模板作为例外处理，不要求表题编号", f"block {index}")
            else:
                table_issue = True
                add(items, "表题位置", "ERROR", "表格前未紧邻表题，或表题格式不合格", f"block {index}", "数据表格前应为：表1  表题；符号说明表前应为“符号说明”")

    if seen_figure and not fig_issue:
        add(items, "图题位置", "PASS", "图题均位于图片下方且位置未发现机器可检问题")
    elif not seen_figure:
        add(items, "图题位置", "INFO", "未检测到图或图片")
    if seen_table and not table_issue:
        add(items, "表题位置", "PASS", "表题均位于表格上方且位置未发现机器可检问题")
    elif not seen_table:
        add(items, "表题位置", "INFO", "未检测到表格")
    return data_tables


def caption_numbers(doc: Document, label: str) -> list[tuple[int, int, str]]:
    pattern = rf"^{label}\s*(\d+)\s+"
    out: list[tuple[int, int, str]] = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        match = re.match(pattern, text)
        if match:
            out.append((idx, int(match.group(1)), text))
    return out


def equation_numbers(doc: Document) -> list[tuple[int, int, str]]:
    out: list[tuple[int, int, str]] = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        match = re.match(r"^\((\d+)\)$", text)
        if match:
            out.append((idx, int(match.group(1)), text))
    return out


def check_number_sequences(doc: Document, items: list[dict[str, str]]) -> None:
    before = len(items)
    groups = [
        ("图", caption_numbers(doc, "图")),
        ("表", caption_numbers(doc, "表")),
        ("公式", equation_numbers(doc)),
    ]
    for label, numbers in groups:
        expected = 1
        for idx, actual, text in numbers:
            if actual != expected:
                add(items, "编号连续性", "ERROR", f"{label}编号不连续：检测到 {actual}，期望 {expected}", f"paragraph {idx}: {text[:80]}", "按全文出现顺序重新编号")
                expected = actual
            expected += 1
    if len(items) == before:
        add(items, "编号连续性", "PASS", "图、表、公式编号连续")


def prior_text(doc: Document, paragraph_index: int) -> str:
    return "\n".join(p.text for p in doc.paragraphs[:paragraph_index])


def has_callout(text: str, label: str, number: int) -> bool:
    if label == "式":
        return f"式({number})" in text or f"式（{number}）" in text
    return f"{label}{number}" in text or f"{label} {number}" in text


def check_callouts(doc: Document, items: list[dict[str, str]]) -> None:
    before = len(items)
    for label, numbers in [("图", caption_numbers(doc, "图")), ("表", caption_numbers(doc, "表"))]:
        for idx, number, text in numbers:
            if not has_callout(prior_text(doc, idx), label, number):
                add(items, "正文呼应", "WARN", f"{label}{number} 前文未检测到呼应文字", f"paragraph {idx}: {text[:80]}", f"正文中先写：如{label}{number}所示 / 见{label}{number}")
    for idx, number, text in equation_numbers(doc):
        if not has_callout(prior_text(doc, idx), "式", number):
            add(items, "正文呼应", "WARN", f"式({number}) 前文未检测到呼应文字", f"paragraph {idx}", f"正文中先写：由式({number})可知")
    if len(items) == before:
        add(items, "正文呼应", "PASS", "正文中已检测到图、表、公式呼应")


def parse_citation_numbers(text: str) -> list[int]:
    numbers: list[int] = []
    for match in re.finditer(r"\[([0-9,\-，、\s]+)\]", text):
        content = match.group(1)
        for part in re.split(r"[,，、\s]+", content):
            if not part:
                continue
            if "-" in part:
                left, right = part.split("-", 1)
                if left.isdigit() and right.isdigit():
                    numbers.extend(range(int(left), int(right) + 1))
            elif part.isdigit():
                numbers.append(int(part))
    return numbers


def check_references(doc: Document, items: list[dict[str, str]]) -> None:
    before = len(items)
    texts = paragraph_texts(doc)
    ref_start = next((i for i, text in enumerate(texts) if text == "参考文献"), None)
    if ref_start is None:
        add(items, "参考文献", "ERROR", "未检测到参考文献标题", fix="添加居中的“参考文献”标题")
        return
    ref_end = next((i for i in range(ref_start + 1, len(texts)) if re.match(r"^附录[A-Z]", texts[i])), len(texts))

    citation_paragraphs = list(enumerate(doc.paragraphs[:ref_start])) + list(enumerate(doc.paragraphs[ref_end:], start=ref_end))
    for para_idx, paragraph in citation_paragraphs:
        for run in paragraph.runs:
            if re.search(r"\[\d+(?:[,，、-]\s*\d+)*\]", run.text):
                if run.font.superscript is not True:
                    add(items, "参考文献", "ERROR", "正文参考文献引用不是上标格式", f"paragraph {para_idx}: {paragraph.text[:80]}", "正文引用应以上标 [1]、[1,2] 形式出现")

    refs: list[tuple[int, int, str]] = []
    for idx in range(ref_start + 1, ref_end):
        text = texts[idx]
        match = re.match(r"^\[(\d+)\](\s+)", text)
        if match:
            refs.append((idx, int(match.group(1)), text))
            if match.group(2) != "　":
                add(items, "参考文献", "WARN", "参考文献序号后不是一字格", f"paragraph {idx}: {text[:80]}", "按《化工进展》规则使用半角序号后接全角空格：[1]　作者...")
        elif re.match(r"^\[\d+\]", text):
            add(items, "参考文献", "ERROR", "参考文献序号后缺少空格", f"paragraph {idx}: {text[:80]}", "序号后空一字格")

    if not refs:
        add(items, "参考文献", "ERROR", "参考文献标题后未检测到编号条目", f"paragraph {ref_start}", "文后文献格式应为：[1] 作者. 题名...")
        return

    for expected, (idx, actual, text) in enumerate(refs, start=1):
        if actual != expected:
            add(items, "参考文献", "ERROR", f"参考文献编号不连续：检测到 [{actual}]，期望 [{expected}]", f"paragraph {idx}: {text[:80]}", "按引用顺序重排并重新编号")
        body = re.sub(r"^\[\d+\]\s*", "", text)
        if not re.search(r"\[[A-Z]+(?:/[A-Z]+)?\]", body):
            add(items, "参考文献", "WARN", "参考文献条目缺少文献类型标识代码", f"paragraph {idx}: {text[:80]}", "按官方规则补 [J]、[M]、[P]、[S]、[D] 等")
        if not body.rstrip().endswith("."):
            add(items, "参考文献", "WARN", "参考文献条目末尾不是半角句点", f"paragraph {idx}: {text[:80]}", "每条文献结束后加半角句点 .")
        if re.search(r"[，。；：．]", body):
            add(items, "参考文献", "WARN", "参考文献中检测到全角标点", f"paragraph {idx}: {text[:80]}", "官方规则要求著录符号使用英文半角符号")
        if re.search(r"[\u4e00-\u9fff]", body):
            next_text = texts[idx + 1].strip() if idx + 1 < len(texts) else ""
            if not next_text or re.match(r"^\[\d+\]\s+", next_text) or re.match(r"^附录[A-Z]", next_text):
                add(items, "参考文献", "WARN", "中文参考文献后未检测到英文对应行", f"paragraph {idx}: {text[:80]}", "中文文献通常需先中文、后英文；若原文献未给英文，人工确认")

    body_text = "\n".join(texts[:ref_start] + texts[ref_end:])
    citation_sequence = parse_citation_numbers(body_text)
    listed = {number for _, number, _ in refs}
    for number in sorted(set(citation_sequence)):
        if number not in listed:
            add(items, "参考文献", "ERROR", f"正文引用了 [{number}]，但文后参考文献缺失", fix="补全文后参考文献或修正引用编号")
    first_seen: list[int] = []
    for number in citation_sequence:
        if number not in first_seen:
            first_seen.append(number)
    if first_seen and first_seen != sorted(first_seen):
        add(items, "参考文献", "ERROR", f"正文首次引用顺序不是递增编号：{first_seen}", fix="参考文献应按正文首次引用顺序排列")
    uncited = sorted(listed - set(citation_sequence))
    if uncited:
        add(items, "参考文献", "WARN", f"文后存在未在正文检测到引用的参考文献：{uncited}", fix="在正文补充引用，或删除未引用文献")

    if len(items) == before:
        add(items, "参考文献", "PASS", "参考文献编号连续，且正文引用顺序未发现问题")


def border_val(cell, edge: str) -> str | None:
    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return None
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        return None
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        return None
    return element.get(qn("w:val"))


def visible_border(cell, edge: str) -> bool:
    value = border_val(cell, edge)
    return value not in {None, "nil", "none"}


def check_three_line_tables(data_tables: list[tuple[int, Any, str]], items: list[dict[str, str]]) -> None:
    before = len(items)
    if not data_tables:
        add(items, "三线表", "WARN", "未检测到带表题的数据表", fix="如论文包含表格，应使用表1  表题 + 三线表")
        return
    for block_idx, table, caption in data_tables:
        if not table.rows:
            add(items, "三线表", "ERROR", "表格没有行", f"block {block_idx}: {caption}", "检查表格数据源")
            continue
        first_row = table.rows[0]
        last_row = table.rows[-1]
        if any(not visible_border(cell, "top") for cell in first_row.cells):
            add(items, "三线表", "WARN", "表头缺少顶线", f"block {block_idx}: {caption}", "三线表应有表头顶线")
        if any(not visible_border(cell, "bottom") for cell in first_row.cells):
            add(items, "三线表", "WARN", "表头缺少表头下线", f"block {block_idx}: {caption}", "三线表应有表头下线")
        if any(not visible_border(cell, "bottom") for cell in last_row.cells):
            add(items, "三线表", "WARN", "表格末行缺少底线", f"block {block_idx}: {caption}", "三线表应有表底线")
        vertical = []
        extra_horizontal = []
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                if visible_border(cell, "left") or visible_border(cell, "right"):
                    vertical.append(row_idx)
                if 0 < row_idx < len(table.rows) - 1 and visible_border(cell, "bottom"):
                    extra_horizontal.append(row_idx)
        if vertical:
            add(items, "三线表", "WARN", "检测到疑似竖线", f"block {block_idx}: {caption}", "三线表通常不保留竖线")
        if extra_horizontal:
            add(items, "三线表", "WARN", "检测到疑似多余横线", f"block {block_idx}: {caption}", "三线表通常只保留顶线、表头下线和底线")
    if len(items) == before:
        add(items, "三线表", "PASS", "数据表疑似符合三线表结构")


def run_size_pt(run) -> float | None:
    return run.font.size.pt if run.font.size is not None else None


def first_visible_run_size(paragraph) -> float | None:
    for run in paragraph.runs:
        if run.text.strip():
            return run_size_pt(run)
    return None


def approx(actual: float | None, expected: float, tol: float = 0.25) -> bool:
    return actual is not None and abs(actual - expected) <= tol


def check_font_sizes(doc: Document, style: dict[str, Any], parts: dict[str, Any], items: list[dict[str, str]]) -> None:
    before = len(items)
    allowed_zh = set(style_at(style, "global_fonts", "chinese_allowed", default=["宋体", "黑体", "仿宋", "楷体"]))
    allowed_aliases = allowed_zh | {"SimSun", "SimHei", "FangSong", "KaiTi", "Times New Roman"}
    xml = parts["document_xml"]
    fonts = set()
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        for match in re.finditer(rf'w:{attr}="([^"]+)"', xml):
            for token in match.group(1).split(";"):
                if token.strip():
                    fonts.add(token.strip())
    unknown = sorted(font for font in fonts if font not in allowed_aliases and not font.startswith("+"))
    if unknown:
        add(items, "字体字号", "WARN", f"检测到未在 STYLE_MAP 允许范围内的字体：{unknown}", fix="中文优先使用宋体、黑体、仿宋、楷体；英文数字用 Times New Roman")

    checks = [
        ("heading_1", r"^\d+\s+\S", 14),
        ("heading_2", r"^\d+\.\d+\s+\S", 10.5),
        ("heading_3", r"^\d+\.\d+\.\d+\s+\S", 10.5),
        ("figure_caption", r"^图\s*\d+\s+", 9),
        ("table_caption", r"^表\s*\d+\s+", 9),
        ("equation", r"^\(\d+\)$", 9),
    ]
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        for name, pattern, fallback in checks:
            if re.match(pattern, text):
                expected = style_at(style, "paragraph_styles", name, "font_size_pt", default=None)
                if expected is None:
                    expected = style_at(style, "paragraph_styles", name, "font_size_pt_observed", default=fallback)
                actual = first_visible_run_size(paragraph)
                if not approx(actual, float(expected)):
                    add(items, "字体字号", "WARN", f"{name} 字号可能不合格：实际 {actual} pt，期望 {expected} pt", f"paragraph {idx}: {text[:80]}", "使用生成器统一直接格式")

    table_size = style_at(style, "paragraph_styles", "table_body", "font_size_pt_observed", default=7.5)
    table_idx = -1
    blocks = list(iter_block_items(doc))
    for block_idx, (kind, table) in enumerate(blocks):
        if kind != "tbl":
            continue
        table_idx += 1
        prev = previous_nonempty(blocks, block_idx)
        if prev and prev[1] == "p" and prev[2].text.strip() == "符号说明":
            continue
        for row in table.rows[:3]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    actual = first_visible_run_size(paragraph)
                    if actual is not None and not approx(actual, float(table_size), tol=0.4) and paragraph.text.strip() != "——":
                        add(items, "字体字号", "WARN", f"表格文字字号可能不合格：实际 {actual} pt，期望约 {table_size} pt", f"table {table_idx}: {paragraph.text[:40]}", "表内文字按 STYLE_MAP 的 table_body 设置")
                        break

    if len(items) == before:
        add(items, "字体字号", "PASS", "常见段落、图题、表题、公式编号和表内字号未发现机器可检问题")


def next_nonempty_index(texts: list[str], start: int, stop: int | None = None) -> int | None:
    stop = len(texts) if stop is None else min(stop, len(texts))
    for idx in range(max(start, 0), stop):
        if texts[idx].strip():
            return idx
    return None


def front_matter_indices(doc: Document, texts: list[str]) -> dict[str, Any]:
    abstract_idx = next((i for i, t in enumerate(texts) if t.startswith("摘要：")), None)
    keywords_idx = next((i for i, t in enumerate(texts) if t.startswith("关键词：")), None)
    classification_idx = next((i for i, t in enumerate(texts) if t.startswith("中图分类号")), None)
    en_abstract_idx = next((i for i, t in enumerate(texts) if re.match(r"^Abstract\s*:", t)), None)
    en_keywords_idx = next((i for i, t in enumerate(texts) if re.match(r"^Keywords\s*:", t)), None)
    doi_idx = next((i for i, t in enumerate(texts[:8]) if t.startswith("DOI")), None)
    article_idx = next_nonempty_index(texts, 0, doi_idx if doi_idx is not None else 2)
    title = find_zh_title(doc, texts, abstract_idx)
    title_idx = title[0] if title else None
    zh_authors_idx = next_nonempty_index(texts, title_idx + 1, abstract_idx) if title_idx is not None and abstract_idx is not None else None
    zh_affiliations_idx = next_nonempty_index(texts, zh_authors_idx + 1, abstract_idx) if zh_authors_idx is not None and abstract_idx is not None else None
    note_indices = []
    if zh_affiliations_idx is not None and abstract_idx is not None:
        note_indices = [i for i in range(zh_affiliations_idx + 1, abstract_idx) if texts[i].strip()]

    en_title_idx = next_nonempty_index(texts, classification_idx + 1, en_abstract_idx) if classification_idx is not None and en_abstract_idx is not None else None
    en_authors_idx = next_nonempty_index(texts, en_title_idx + 1, en_abstract_idx) if en_title_idx is not None and en_abstract_idx is not None else None
    en_affiliations_idx = next_nonempty_index(texts, en_authors_idx + 1, en_abstract_idx) if en_authors_idx is not None and en_abstract_idx is not None else None

    return {
        "article_type": article_idx,
        "doi": doi_idx,
        "zh_title": title_idx,
        "zh_authors": zh_authors_idx,
        "zh_affiliations": zh_affiliations_idx,
        "front_matter_notes": note_indices,
        "zh_abstract": abstract_idx,
        "zh_keywords": keywords_idx,
        "classification_line": classification_idx,
        "en_title": en_title_idx,
        "en_authors": en_authors_idx,
        "en_affiliations": en_affiliations_idx,
        "en_abstract": en_abstract_idx,
        "en_keywords": en_keywords_idx,
    }


def check_front_matter_structure(doc: Document, style: dict[str, Any], parts: dict[str, Any], items: list[dict[str, str]]) -> dict[str, Any]:
    before = len(items)
    paragraphs = body_paragraph_elements(parts)
    texts = [xml_text(p) for p in paragraphs]
    indices = front_matter_indices(doc, texts)
    expected_order = style_at(style, "strict_layout", "front_matter_sequence", default=[]) or []
    flattened: list[tuple[str, int]] = []
    for name in expected_order:
        value = indices.get(name)
        if isinstance(value, list):
            flattened.extend((name, idx) for idx in value)
        elif value is not None:
            flattened.append((name, int(value)))
        elif name != "front_matter_notes":
            add(items, "前置换行结构", "ERROR", f"未检测到前置部分：{name}", fix="按模板将题名、作者、摘要、关键词等分别放在独立段落")
    positions = [idx for _, idx in flattened]
    if positions != sorted(positions):
        add(items, "前置换行结构", "ERROR", "前置部分顺序与模板不一致", str(flattened), "按文章类型、DOI、中文题名、作者单位、摘要关键词、中图分类号、英文信息顺序排布")
    seen_positions: dict[int, list[str]] = {}
    for name, idx in flattened:
        seen_positions.setdefault(idx, []).append(name)
    merged = {idx: names for idx, names in seen_positions.items() if len(names) > 1}
    if merged:
        add(items, "前置换行结构", "ERROR", f"检测到多个前置内容共用同一段落：{merged}", fix="不同内容必须按模板独立成段，不要用软换行或同段拼接")
    if len(items) == before:
        add(items, "前置换行结构", "PASS", "前置部分独立段落、顺序和换行结构未发现机器可检问题")
    return indices


def check_run_against_config(
    run: etree._Element | None,
    cfg: dict[str, Any],
    style: dict[str, Any],
    style_name: str,
    role: str,
    idx: int,
    text: str,
    items: list[dict[str, str]],
) -> None:
    if run is None or not cfg:
        return
    r_fonts = rpr_child_attrs(run, "rFonts")
    r_sz = rpr_child_attrs(run, "sz")
    r_szcs = rpr_child_attrs(run, "szCs")
    r_spacing = rpr_child_attrs(run, "spacing")
    r_bold = rpr_child_attrs(run, "b")
    prefix = f"{style_name}.{role}"

    expected_font_zh = cfg.get("font_zh")
    if expected_font_zh and font_token(r_fonts.get("eastAsia")) != font_token(expected_font_zh):
        add(items, "细排版", "WARN", f"{prefix} 中文字体不合格：实际 {r_fonts.get('eastAsia')}，期望 {expected_font_zh}", f"paragraph {idx}: {text[:80]}", "按 strict_layout.styles 的 run 级规则写入 w:rFonts/@w:eastAsia")

    expected_latin = cfg.get("font_latin")
    if expected_latin:
        for attr in ("ascii", "hAnsi"):
            if font_token(r_fonts.get(attr)) != font_token(expected_latin):
                add(items, "细排版", "WARN", f"{prefix} {attr} 字体不合格：实际 {r_fonts.get(attr)}，期望 {expected_latin}", f"paragraph {idx}: {text[:80]}", "字母和数字应使用 Times New Roman")

    half_tol = int(style_at(style, "strict_layout", "tolerance", "half_points", default=1))
    for label, actual_attrs, expected in [("sz", r_sz, cfg.get("sz")), ("szCs", r_szcs, cfg.get("szCs"))]:
        if expected is None:
            continue
        actual = raw_int(actual_attrs.get("val"))
        target = raw_int(str(expected))
        if actual is None or target is None or abs(actual - target) > half_tol:
            add(items, "细排版", "WARN", f"{prefix} {label} 字号不合格：实际 {actual_attrs.get('val')}，期望 {expected}", f"paragraph {idx}: {text[:80]}", "按模板半磅值写入 w:sz/w:szCs")

    expected_char_spacing = cfg.get("character_spacing")
    if expected_char_spacing is not None and str(r_spacing.get("val")) != str(expected_char_spacing):
        add(items, "细排版", "WARN", f"{prefix} 字符间距不合格：实际 {r_spacing.get('val')}，期望 {expected_char_spacing}", f"paragraph {idx}: {text[:80]}", "按模板 raw w:rPr/w:spacing 设置")

    expected_bold = cfg.get("bold")
    if expected_bold is True and (not rpr_has_child(run, "b") or r_bold.get("val", "1") in {"0", "false", "False"}):
        add(items, "细排版", "WARN", f"{prefix} 应加粗但未检测到加粗", f"paragraph {idx}: {text[:80]}", "按模板设置 w:b")
    if expected_bold is False and rpr_has_child(run, "b") and r_bold.get("val", "1") not in {"0", "false", "False"}:
        add(items, "细排版", "WARN", f"{prefix} 不应加粗但检测到加粗", f"paragraph {idx}: {text[:80]}", "移除该 run 的 w:b")


def check_style_on_paragraph(paragraphs: list[etree._Element], idx: int, style: dict[str, Any], style_name: str, items: list[dict[str, str]]) -> None:
    if idx < 0 or idx >= len(paragraphs):
        return
    st = style_at(style, "strict_layout", "styles", style_name, default={}) or {}
    if not st:
        return
    tol = style_at(style, "strict_layout", "tolerance", "twips", default=8)
    paragraph = paragraphs[idx]
    text = xml_text(paragraph)
    p_jc = ppr_child_attrs(paragraph, "jc")
    p_ind = ppr_child_attrs(paragraph, "ind")
    p_spacing = ppr_child_attrs(paragraph, "spacing")
    run = first_text_run(paragraph)
    r_fonts = rpr_child_attrs(run, "rFonts")
    r_sz = rpr_child_attrs(run, "sz")
    r_szcs = rpr_child_attrs(run, "szCs")
    r_spacing = rpr_child_attrs(run, "spacing")
    r_bold = rpr_child_attrs(run, "b")

    expected_jc = st.get("jc")
    if expected_jc is not None and p_jc.get("val") != str(expected_jc):
        add(items, "细排版", "WARN", f"{style_name} 对齐不合格：实际 {p_jc.get('val')}，期望 {expected_jc}", f"paragraph {idx}: {text[:80]}", "按 strict_layout.styles 写入 w:jc")

    expected_ind = dict(st.get("ind") or {})
    for key in ("firstLine", "hanging", "left", "right"):
        if st.get(key) is not None:
            expected_ind[key] = st[key]
    for key, expected in expected_ind.items():
        if not near_twips(p_ind.get(key), expected, int(tol)):
            add(items, "细排版", "WARN", f"{style_name} 缩进 {key} 不合格：实际 {p_ind.get(key)}，期望 {expected}", f"paragraph {idx}: {text[:80]}", "按模板 raw w:ind 设置，不使用近似 cm")

    for key in ("line", "lineRule", "before", "after"):
        expected = st.get(key)
        if expected is None and key in {"line", "lineRule"}:
            expected = style_at(style, "strict_layout", "paragraph_defaults", key, default=None)
        if expected is None:
            continue
        actual = p_spacing.get(key)
        if str(actual) != str(expected):
            add(items, "细排版", "WARN", f"{style_name} 段落间距/行距 {key} 不合格：实际 {actual}，期望 {expected}", f"paragraph {idx}: {text[:80]}", "按模板 raw w:spacing 设置")

    label_cfg = st.get("label_run") if isinstance(st.get("label_run"), dict) else None
    body_cfg = st.get("body_run") if isinstance(st.get("body_run"), dict) else None
    if label_cfg or body_cfg:
        runs = text_runs(paragraph)
        labels = [str(label) for label in (st.get("label_texts") or [])]
        if labels:
            for item_run in runs:
                run_text = xml_text(item_run)
                if any(run_text.startswith(label) for label in labels):
                    check_run_against_config(item_run, label_cfg or {}, style, style_name, "label_run", idx, text, items)
                elif run_text.strip():
                    check_run_against_config(item_run, body_cfg or {}, style, style_name, "body_run", idx, text, items)
        else:
            if runs:
                check_run_against_config(runs[0], label_cfg or st, style, style_name, "label_run", idx, text, items)
            body_run = next((candidate for candidate in runs[1:] if xml_text(candidate).strip()), None)
            if body_run is not None:
                check_run_against_config(body_run, body_cfg or st, style, style_name, "body_run", idx, text, items)
            elif body_cfg:
                add(items, "细排版", "WARN", f"{style_name} 未检测到正文内容 run", f"paragraph {idx}: {text[:80]}", "label 和正文应拆成独立 run")
        return

    expected_font_zh = st.get("font_zh")
    if expected_font_zh and font_token(r_fonts.get("eastAsia")) != font_token(expected_font_zh):
        add(items, "细排版", "WARN", f"{style_name} 中文字体不合格：实际 {r_fonts.get('eastAsia')}，期望 {expected_font_zh}", f"paragraph {idx}: {text[:80]}", "按 strict_layout.styles 写入 w:rFonts/@w:eastAsia")
    expected_latin = st.get("font_latin")
    if expected_latin:
        for attr in ("ascii", "hAnsi"):
            if font_token(r_fonts.get(attr)) != font_token(expected_latin):
                add(items, "细排版", "WARN", f"{style_name} {attr} 字体不合格：实际 {r_fonts.get(attr)}，期望 {expected_latin}", f"paragraph {idx}: {text[:80]}", "字母和数字应使用 Times New Roman")

    half_tol = int(style_at(style, "strict_layout", "tolerance", "half_points", default=1))
    for label, actual_attrs, expected in [("sz", r_sz, st.get("sz")), ("szCs", r_szcs, st.get("szCs"))]:
        if expected is None:
            continue
        actual = raw_int(actual_attrs.get("val"))
        target = raw_int(str(expected))
        if actual is None or target is None or abs(actual - target) > half_tol:
            add(items, "细排版", "WARN", f"{style_name} {label} 字号不合格：实际 {actual_attrs.get('val')}，期望 {expected}", f"paragraph {idx}: {text[:80]}", "按模板半磅值写入 w:sz/w:szCs")

    expected_char_spacing = st.get("character_spacing")
    if expected_char_spacing is not None and str(r_spacing.get("val")) != str(expected_char_spacing):
        add(items, "细排版", "WARN", f"{style_name} 字符间距不合格：实际 {r_spacing.get('val')}，期望 {expected_char_spacing}", f"paragraph {idx}: {text[:80]}", "按模板 raw w:rPr/w:spacing 设置")

    if st.get("bold") is True and (not rpr_has_child(run, "b") or r_bold.get("val", "1") in {"0", "false", "False"}):
        add(items, "细排版", "WARN", f"{style_name} 应加粗但未检测到加粗", f"paragraph {idx}: {text[:80]}", "按模板设置 w:b")


def strict_style_for_text(text: str, ref_start: int | None) -> str | None:
    if re.match(r"^\d+\.\d+\.\d+\s+\S", text):
        return "heading_3"
    if re.match(r"^\d+\.\d+\s+\S", text):
        return "heading_2"
    if re.match(r"^\d+\s+\S", text):
        return "heading_1"
    if re.match(r"^附录[A-Z](?:\s+|$)", text):
        return "appendix"
    if re.match(r"^图\s*\d+\s+", text):
        return "figure_caption"
    if re.match(r"^表\s*\d+\s+", text):
        return "table_caption"
    if re.match(r"^\(\d+\)$", text):
        return "equation_number"
    if text in {"参考文献", "符号说明"}:
        return "table_caption"
    return "references" if ref_start is not None else None


def check_strict_layout(doc: Document, style: dict[str, Any], parts: dict[str, Any], items: list[dict[str, str]]) -> None:
    if not style_at(style, "strict_layout", "enabled", default=False):
        return
    before = len(items)
    paragraphs = body_paragraph_elements(parts)
    texts = [xml_text(p) for p in paragraphs]
    front = check_front_matter_structure(doc, style, parts, items)
    checked: set[int] = set()
    for name, value in front.items():
        if isinstance(value, list):
            for idx in value:
                check_style_on_paragraph(paragraphs, idx, style, "front_matter_note", items)
                checked.add(idx)
        elif value is not None:
            check_style_on_paragraph(paragraphs, int(value), style, name, items)
            checked.add(int(value))

    ref_title_idx = next((i for i, text in enumerate(texts) if text == "参考文献"), None)
    appendix_start_idx = next((i for i, text in enumerate(texts) if re.match(r"^附录[A-Z](?:\s+|$)", text)), None)
    en_keywords_idx = front.get("en_keywords")
    for idx, text in enumerate(texts):
        if idx in checked or not text:
            continue
        ref_context = (
            ref_title_idx
            if ref_title_idx is not None and idx > ref_title_idx and (appendix_start_idx is None or idx < appendix_start_idx)
            else None
        )
        name = strict_style_for_text(text, ref_context)
        if name is None and idx > 0 and re.match(r"^图\s*\d+\s+", texts[idx - 1]):
            name = "figure_note"
        if name is None and isinstance(en_keywords_idx, int) and idx > en_keywords_idx and (ref_title_idx is None or idx < ref_title_idx):
            if re.match(r"^附录[A-Z](?:\s+|$)", text):
                name = "appendix"
            elif not paragraph_has_image(doc.paragraphs[idx]) and not paragraph_has_math(doc.paragraphs[idx]):
                name = "body"
        if name is None and appendix_start_idx is not None and idx >= appendix_start_idx:
            name = "appendix" if re.match(r"^附录[A-Z](?:\s+|$)", text) else "body"
        if name is not None:
            check_style_on_paragraph(paragraphs, idx, style, name, items)

    if not any(item["category"] == "细排版" for item in items[before:]):
        add(items, "细排版", "PASS", "段落对齐、缩进、行距、字体、字号、复杂文种字号和字符间距均符合 strict_layout")


def check_table_micro(data_tables: list[tuple[int, Any, str]], style: dict[str, Any], items: list[dict[str, str]]) -> None:
    before = len(items)
    expected = style_at(style, "strict_layout", "tables", default={}) or {}
    if not expected:
        return
    tblw_expected = expected.get("tblW") or {}
    cellmar_expected = expected.get("cellMar") or {}
    borders_expected = expected.get("borders") or {}
    for block_idx, table, caption in data_tables:
        tbl_pr = table._tbl.tblPr
        tbl_w = local_attrs(tbl_pr.find(qn("w:tblW")) if tbl_pr is not None else None)
        for key, value in tblw_expected.items():
            if str(tbl_w.get(key)) != str(value):
                add(items, "表格细节", "WARN", f"表宽 {key} 不合格：实际 {tbl_w.get(key)}，期望 {value}", f"block {block_idx}: {caption}", "按模板设置 w:tblW")
        cell_mar = tbl_pr.find(qn("w:tblCellMar")) if tbl_pr is not None else None
        for edge, value in cellmar_expected.items():
            edge_attrs = local_attrs(cell_mar.find(qn(f"w:{edge}")) if cell_mar is not None else None)
            if str(edge_attrs.get("w")) != str(value):
                add(items, "表格细节", "WARN", f"单元格边距 {edge} 不合格：实际 {edge_attrs.get('w')}，期望 {value}", f"block {block_idx}: {caption}", "按模板设置 tblCellMar")
        if not table.rows:
            continue
        samples = [
            ("表头顶线", table.rows[0].cells[0], "top", borders_expected.get("header_top")),
            ("表头下线", table.rows[0].cells[0], "bottom", borders_expected.get("header_bottom")),
            ("表底线", table.rows[-1].cells[0], "bottom", borders_expected.get("table_bottom")),
        ]
        for label, cell, edge, exp in samples:
            if not exp:
                continue
            tc_pr = cell._tc.tcPr
            tc_borders = tc_pr.find(qn("w:tcBorders")) if tc_pr is not None else None
            actual = local_attrs(tc_borders.find(qn(f"w:{edge}")) if tc_borders is not None else None)
            for key in ("val", "sz", "color"):
                if str(actual.get(key)) != str(exp.get(key)):
                    add(items, "表格细节", "WARN", f"{label} {key} 不合格：实际 {actual.get(key)}，期望 {exp.get(key)}", f"block {block_idx}: {caption}", "按模板三线表线宽写入 tcBorders")
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                tc_pr = cell._tc.tcPr
                tc_borders = tc_pr.find(qn("w:tcBorders")) if tc_pr is not None else None
                for edge in ("left", "right", "insideV"):
                    if local_attrs(tc_borders.find(qn(f"w:{edge}")) if tc_borders is not None else None).get("val") not in {None, "", "nil", "none"}:
                        add(items, "表格细节", "WARN", "检测到模板三线表不需要的竖线", f"block {block_idx}, row {row_idx}, col {col_idx}: {caption}", "删除竖线，只保留顶线、表头下线和底线")
    if len(items) == before:
        add(items, "表格细节", "PASS", "表宽、单元格边距和三线表 raw 边框属性符合 strict_layout")


def check_images_and_equations(doc: Document, parts: dict[str, Any], items: list[dict[str, str]]) -> None:
    xml = parts["document_xml"]
    if "<wp:anchor" in xml:
        add(items, "图片嵌入", "ERROR", "检测到浮动图片 wp:anchor", fix="全部图片改为嵌入型")
    else:
        add(items, "图片嵌入", "PASS", "未检测到浮动图片；图片为嵌入型或无图片")

    eq_nums = equation_numbers(doc)
    if eq_nums:
        if "<m:oMath" not in xml and "<m:oMathPara" not in xml:
            add(items, "公式", "WARN", "检测到公式编号，但未检测到 Office Math/OMML 对象", fix="正式稿应使用 MathType 或可验证的公式对象，不要只用普通文本")
        else:
            add(items, "公式", "PASS", "检测到公式编号和 OMML 公式对象")
    else:
        add(items, "公式", "INFO", "未检测到公式编号")


def check_appendices(doc: Document, items: list[dict[str, str]]) -> None:
    before = len(items)
    letters: list[str] = []
    for idx, text in enumerate(paragraph_texts(doc)):
        match = re.match(r"^附录([A-Z])(?:\s+|$)", text)
        if match:
            letters.append(match.group(1))
            expected = chr(ord("A") + len(letters) - 1)
            if match.group(1) != expected:
                add(items, "附录", "WARN", f"附录编号不连续：检测到 {match.group(1)}，期望 {expected}", f"paragraph {idx}: {text[:80]}", "按附录A、附录B顺序编号")
    if not letters:
        add(items, "附录", "INFO", "未检测到附录")
    elif len(items) == before:
        add(items, "附录", "PASS", "附录编号未发现机器可检问题")


def check_notes(parts: dict[str, Any], items: list[dict[str, str]]) -> None:
    footnotes = real_note_ids(parts["parts"].get("word/footnotes.xml", ""), "footnote")
    endnotes = real_note_ids(parts["parts"].get("word/endnotes.xml", ""), "endnote")
    if footnotes or endnotes:
        add(items, "注释", "ERROR", f"检测到脚注/尾注编号：footnotes={footnotes}, endnotes={endnotes}", fix="提交稿不要保留 Word 脚注/尾注")
    else:
        add(items, "注释", "PASS", "未检测到 Word 脚注或尾注；图表注应为普通文本")


def compact_len(text: str) -> int:
    return len(re.sub(r"\s+", "", text or ""))


def split_keywords(text: str) -> list[str]:
    body = re.sub(r"^\s*(关键词|Keywords)\s*[:：]\s*", "", text, flags=re.I)
    return [part.strip() for part in re.split(r"[；;,，]", body) if part.strip()]


def find_zh_title(doc: Document, texts: list[str], abstract_idx: int | None) -> tuple[int, str] | None:
    stop = abstract_idx if abstract_idx is not None else min(len(texts), 12)
    candidates: list[tuple[int, str, int]] = []
    for idx, paragraph in enumerate(doc.paragraphs[:stop]):
        text = paragraph.text.strip()
        if not text or len(text) > 60:
            continue
        if any(text.startswith(prefix) for prefix in ["DOI", "文章类型", "摘要", "关键词", "中图分类号", "Abstract", "Keywords"]):
            continue
        if text.startswith("（") or "，" in text or "," in text or "；" in text or ":" in text or "：" in text:
            continue
        centered = paragraph.alignment == 1 or "CENTER" in str(paragraph.alignment).upper()
        cjk = bool(re.search(r"[\u4e00-\u9fff]", text))
        if centered and cjk:
            candidates.append((idx, text, compact_len(text)))
    if not candidates:
        return None
    return max(candidates, key=lambda item: item[2])[:2]


def check_title_abstract_keywords(doc: Document, style: dict[str, Any], items: list[dict[str, str]]) -> None:
    before = len(items)
    rules = style_at(style, "content_rules", default={}) or {}
    texts = paragraph_texts(doc)
    abstract_idx = next((i for i, t in enumerate(texts) if t.startswith("摘要：")), None)
    keywords_idx = next((i for i, t in enumerate(texts) if t.startswith("关键词：")), None)
    en_abstract_idx = next((i for i, t in enumerate(texts) if re.match(r"^Abstract\s*:", t)), None)
    en_keywords_idx = next((i for i, t in enumerate(texts) if re.match(r"^Keywords\s*:", t)), None)

    title = find_zh_title(doc, texts, abstract_idx)
    if title:
        title_idx, title_text = title
        max_chars = int(style_at(rules, "zh_title", "max_chars", default=20))
        if compact_len(title_text) > max_chars:
            add(items, "标题摘要关键词", "WARN", f"中文题名超过 {max_chars} 字", f"paragraph {title_idx}: {title_text}", "题名应精炼，通常不超过 20 字")
        avoid = style_at(rules, "zh_title", "avoid_words", default=[]) or []
        bad = [word for word in avoid if word in title_text]
        if bad:
            add(items, "标题摘要关键词", "WARN", f"中文题名含弱信息词：{bad}", f"paragraph {title_idx}: {title_text}", "按模板建议去掉“研究”“实验”“综述”等泛化词")
    else:
        add(items, "标题摘要关键词", "WARN", "未能可靠识别中文题名", fix="检查首页题名段落是否居中且位于摘要前")

    if abstract_idx is None:
        add(items, "标题摘要关键词", "ERROR", "未检测到中文摘要", fix="添加“摘要：”段落")
    else:
        abstract = re.sub(r"^摘要：", "", texts[abstract_idx]).strip()
        min_chars = int(style_at(rules, "abstract_zh", "min_chars", default=300))
        if compact_len(abstract) < min_chars:
            add(items, "标题摘要关键词", "WARN", f"中文摘要不足 {min_chars} 字：当前约 {compact_len(abstract)} 字", f"paragraph {abstract_idx}", "摘要需写出研究问题、过程和方法、结果，并进行定性定量表述")
        avoid_starts = style_at(rules, "abstract_zh", "avoid_starts", default=["本文", "作者"]) or []
        if any(abstract.startswith(word) for word in avoid_starts):
            add(items, "语言风格", "WARN", "摘要以“本文/作者”等作主语", f"paragraph {abstract_idx}", "摘要不以“本文”“作者”等作主语，直接陈述研究内容")
        if re.search(r"\[\d+(?:[,，、-]\d+)*\]", abstract):
            add(items, "标题摘要关键词", "WARN", "摘要中检测到参考文献编号", f"paragraph {abstract_idx}", "摘要中一般不出现参考文献序号")
        if "=" in abstract or re.search(r"式\s*[（(]\d+[）)]", abstract):
            add(items, "标题摘要关键词", "WARN", "摘要中疑似出现公式或公式引用", f"paragraph {abstract_idx}", "摘要中一般不出现公式")

    if keywords_idx is None:
        add(items, "标题摘要关键词", "ERROR", "未检测到中文关键词", fix="添加“关键词：”段落")
    else:
        keywords = split_keywords(texts[keywords_idx])
        min_count = int(style_at(rules, "keywords_zh", "min_count", default=3))
        max_count = int(style_at(rules, "keywords_zh", "max_count", default=6))
        preferred_max = int(style_at(rules, "keywords_zh", "preferred_max_count", default=max_count))
        if not (min_count <= len(keywords) <= max_count):
            add(items, "标题摘要关键词", "WARN", f"中文关键词数量为 {len(keywords)}，期望 {min_count}-{max_count} 个", f"paragraph {keywords_idx}: {texts[keywords_idx]}", "按征稿简则控制关键词数量")
        elif len(keywords) > preferred_max:
            add(items, "标题摘要关键词", "INFO", f"中文关键词数量为 {len(keywords)}，修改要求提示一般 {min_count}-{preferred_max} 个", f"paragraph {keywords_idx}")
        forbidden = style_at(rules, "keywords_zh", "forbidden", default=["进展", "综述", "应用", "研究"]) or []
        bad_keywords = [kw for kw in keywords if kw in forbidden]
        if bad_keywords:
            add(items, "标题摘要关键词", "WARN", f"关键词使用了禁用泛词：{bad_keywords}", f"paragraph {keywords_idx}", "“进展”“综述”“应用”“研究”等不作为关键词")
        if "；" not in texts[keywords_idx] and len(keywords) > 1:
            add(items, "标题摘要关键词", "WARN", "中文关键词未使用中文分号分隔", f"paragraph {keywords_idx}", "中文关键词之间用“；”")
        min_standard = int(style_at(rules, "keywords_zh", "min_standard_count", default=0) or 0)
        if min_standard > 0:
            standard_keywords, standard_source, standard_status = load_standard_keywords(style)
            if standard_keywords:
                matched = [kw for kw in keywords if kw in standard_keywords]
                if len(matched) < min_standard:
                    missing = [kw for kw in keywords if kw not in standard_keywords]
                    add(
                        items,
                        "标题摘要关键词",
                        "WARN",
                        f"标准关键词命中 {len(matched)} 个，少于要求的 {min_standard} 个",
                        f"paragraph {keywords_idx}: {texts[keywords_idx]}",
                        f"从标准关键词库补足至少 {min_standard} 个；当前未命中：{missing}",
                    )
            else:
                add(
                    items,
                    "标题摘要关键词",
                    "INFO",
                    f"未加载标准关键词库：{standard_status}",
                    standard_source,
                    "检查 STYLE_MAP.json 中 standard_keyword_file 是否指向本地 JSON/TXT 词库",
                )

    if en_abstract_idx is None:
        add(items, "标题摘要关键词", "WARN", "未检测到英文摘要", fix="添加 Abstract: 段落")
    if en_keywords_idx is None:
        add(items, "标题摘要关键词", "WARN", "未检测到英文关键词", fix="添加 Keywords: 段落")

    if len(items) == before:
        add(items, "标题摘要关键词", "PASS", "题名、摘要和关键词未发现机器可检语言问题")


def check_language_units_conclusion(doc: Document, style: dict[str, Any], items: list[dict[str, str]]) -> None:
    rules = style_at(style, "content_rules", default={}) or {}
    texts = paragraph_texts(doc)
    ref_start = next((i for i, text in enumerate(texts) if text == "参考文献"), len(texts))
    body_texts = texts[:ref_start]

    before_lang = len(items)
    first_person = style_at(rules, "body", "avoid_first_person", default=["我们", "本人", "本作者"]) or []
    for idx, text in enumerate(body_texts):
        if not text:
            continue
        hits = [word for word in first_person if word in text]
        if hits:
            add(items, "语言风格", "WARN", f"检测到第一人称或作者自称：{hits}", f"paragraph {idx}: {text[:80]}", "按模板要求尽量不用“我们”“本人”等第一人称")
        if re.match(r"^\d+(?:\.\d+)*\s+引言|^\d+(?:\.\d+)*\s+前言", text):
            add(items, "语言风格", "ERROR", "引言/前言被编号", f"paragraph {idx}: {text[:80]}", "引言不写标题或不排序号")
    if len(items) == before_lang:
        add(items, "语言风格", "PASS", "未发现第一人称、编号引言等机器可检语言问题")

    before_units = len(items)
    legacy_units = style_at(rules, "body", "banned_legacy_units", default=[]) or []
    for idx, text in enumerate(body_texts):
        for unit in legacy_units:
            if unit == "in":
                pattern = r"\d(?:\s|~|～|-)*in(?![A-Za-z])"
            else:
                pattern = rf"\d(?:\s|~|～|-)*{re.escape(unit)}(?![A-Za-z])"
            if re.search(pattern, text):
                add(items, "单位符号", "WARN", f"检测到需换算或确认的非标准单位：{unit}", f"paragraph {idx}: {text[:80]}", "按修改要求换成国际制/法定计量单位")
    if len(items) == before_units:
        add(items, "单位符号", "PASS", "未检测到 ppm、cal、atm、mmHg、in、bar、kgf 等需重点确认的单位")

    before_conclusion = len(items)
    conclusion_indices = [i for i, text in enumerate(texts[:ref_start]) if re.match(r"^\d+\s+结论|^\d+\s+结语|^结论$|^结语$", text)]
    if not conclusion_indices:
        add(items, "结论结语", "WARN", "未检测到结论/结语章节", fix="研究性论文用“结论”，综述性论文一般用“结语”")
    else:
        start = conclusion_indices[-1]
        end = ref_start
        conclusion_text = "\n".join(texts[start:end])
        if "结论" in texts[start]:
            for word in ["存在问题", "不足", "前瞻", "展望"]:
                if word in conclusion_text:
                    break
            else:
                add(items, "结论结语", "WARN", "研究性论文结论中未检测到存在问题/不足/展望/前瞻性表述", f"paragraph {start}: {texts[start]}", "修改基本要求提示研究性论文需给出结论以及存在的问题和前瞻性")
        if "结语" in texts[start]:
            for word in ["发展趋势", "存在问题", "前景", "展望"]:
                if word in conclusion_text:
                    break
            else:
                add(items, "结论结语", "WARN", "综述性论文结语中未检测到发展趋势/存在问题/前景展望", f"paragraph {start}: {texts[start]}", "综述性论文需给出具体研究发展趋势、存在问题、前景展望")
    if len(items) == before_conclusion:
        add(items, "结论结语", "PASS", "结论/结语结构未发现机器可检问题")


def write_report(items: list[dict[str, str]], out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    categories = MANDATORY_CATEGORIES + [
        c for c in EXTRA_CATEGORIES + CONTENT_CATEGORIES if any(i["category"] == c for i in items)
    ]
    counts = {"ERROR": 0, "WARN": 0, "INFO": 0, "PASS": 0}
    for item in items:
        counts[item["status"]] = counts.get(item["status"], 0) + 1

    lines = [
        "# DOCX 格式检查报告",
        "",
        f"ERROR: {counts.get('ERROR', 0)}，WARN: {counts.get('WARN', 0)}，INFO: {counts.get('INFO', 0)}",
        "",
        "## 检查项总览",
        "",
        "| 检查项 | 结果 |",
        "|---|---|",
    ]
    for category in categories:
        lines.append(f"| {category} | {summarize_category(items, category)} |")
    lines.append("")

    for category in categories:
        lines.append(f"## {category}")
        category_items = [item for item in items if item["category"] == category]
        if not category_items:
            lines.append("- INFO：未执行该检查项。")
            lines.append("")
            continue
        for item in category_items:
            lines.append(f"- {item['status']}：{item['message']}")
            if item.get("location"):
                lines.append(f"  - 位置：{item['location']}")
            if item.get("fix"):
                lines.append(f"  - 建议：{item['fix']}")
        lines.append("")
    lines.append("## 说明")
    lines.append("- 本报告覆盖结构、编号、禁用 Word 特性、页面设置、常见字体字号、摘要关键词、语言风格、单位符号和结论/结语。")
    lines.append("- 图片内部线宽、坐标轴、白边裁剪、真实图像质量和 MathType 兼容性仍需渲染后复核或人工确认。")
    out.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--style", type=Path, default=Path("STYLE_MAP.json"))
    parser.add_argument("--out", type=Path, default=Path("format_report.md"))
    args = parser.parse_args()

    style = load_style(args.style)
    parts = read_docx_parts(args.docx)
    doc = Document(args.docx)
    items: list[dict[str, str]] = []

    check_page(doc, style, parts, items)
    check_forbidden_features(parts, items)
    check_headings(doc, items)
    data_tables = check_caption_positions(doc, items)
    check_number_sequences(doc, items)
    check_callouts(doc, items)
    check_references(doc, items)
    check_three_line_tables(data_tables, items)
    check_table_micro(data_tables, style, items)
    check_font_sizes(doc, style, parts, items)
    check_strict_layout(doc, style, parts, items)
    check_images_and_equations(doc, parts, items)
    check_appendices(doc, items)
    check_notes(parts, items)
    check_title_abstract_keywords(doc, style, items)
    check_language_units_conclusion(doc, style, items)
    write_report(items, args.out)

    errors = sum(1 for item in items if item["status"] == "ERROR")
    warnings = sum(1 for item in items if item["status"] == "WARN")
    print(f"Wrote {args.out} with {errors} errors and {warnings} warnings")


if __name__ == "__main__":
    main()
