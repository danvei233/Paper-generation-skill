#!/usr/bin/env python3
"""Extract useful style/layout facts from a DOCX template.

This script is intentionally conservative. It does not claim every observed
format is a hard rule. It emits a JSON report that Codex/humans can compare
against FORMAT_SPEC.md and STYLE_MAP.json.
"""
from __future__ import annotations

import argparse, json, re, zipfile
from pathlib import Path
from typing import Any
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

NS = {"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

def cm(x):
    return round(x.cm, 4) if x is not None else None

def run_info(run):
    f = run.font
    return {
        "text": run.text[:80],
        "font": f.name,
        "size_pt": f.size.pt if f.size else None,
        "bold": f.bold,
        "italic": f.italic,
        "underline": bool(f.underline),
        "color": str(f.color.rgb) if f.color and f.color.rgb else None,
    }

def paragraph_info(i, p):
    pf = p.paragraph_format
    return {
        "index": i,
        "text": p.text[:250],
        "style": p.style.name if p.style else None,
        "alignment": str(p.alignment),
        "left_indent_cm": cm(pf.left_indent),
        "right_indent_cm": cm(pf.right_indent),
        "first_line_indent_cm": cm(pf.first_line_indent),
        "space_before_pt": pf.space_before.pt if pf.space_before else None,
        "space_after_pt": pf.space_after.pt if pf.space_after else None,
        "line_spacing": str(pf.line_spacing),
        "runs": [run_info(r) for r in p.runs if r.text.strip()][:10],
    }

def attr_map(element):
    if element is None:
        return None
    return {k.split("}")[-1]: v for k, v in element.attrib.items()}

def local_attrs(element):
    if element is None:
        return {}
    return {etree.QName(k).localname: v for k, v in element.attrib.items()}

def child_attrs(element, child_name: str):
    if element is None:
        return {}
    child = element.find(f"w:{child_name}", namespaces=NS)
    return local_attrs(child)

def xml_text(element) -> str:
    if element is None:
        return ""
    return "".join(element.xpath(".//w:t/text()", namespaces=NS)).strip()

def first_text_run(paragraph):
    if paragraph is None:
        return None
    runs = paragraph.xpath("./w:r[w:t]", namespaces=NS)
    return runs[0] if runs else None

def guess_strict_style(text: str) -> str | None:
    if text.startswith("文章类型："):
        return "article_type"
    if text.startswith("DOI"):
        return "doi"
    if text.startswith("摘要："):
        return "zh_abstract"
    if text.startswith("关键词："):
        return "zh_keywords"
    if text.startswith("中图分类号"):
        return "classification_line"
    if text.startswith("Abstract:"):
        return "en_abstract"
    if text.startswith("Keywords:"):
        return "en_keywords"
    if text == "符号说明":
        return "table_caption"
    if text == "参考文献":
        return "table_caption"
    if re.match(r"^附录[A-Z](?:\s+|$)", text):
        return "appendix"
    if re.match(r"^\d+\.\d+\.\d+\s+\S", text):
        return "heading_3"
    if re.match(r"^\d+\.\d+\s+\S", text):
        return "heading_2"
    if re.match(r"^\d+\s+\S", text):
        return "heading_1"
    if re.match(r"^图\s*\d+\s+", text):
        return "figure_caption"
    if re.match(r"^表\s*\d+\s+", text):
        return "table_caption"
    if re.match(r"^\(\d+\)$", text):
        return "equation_number"
    return None

def raw_paragraph_observation(index: int, paragraph) -> dict[str, Any]:
    ppr = paragraph.find("w:pPr", namespaces=NS)
    run = first_text_run(paragraph)
    rpr = run.find("w:rPr", namespaces=NS) if run is not None else None
    first_runs = []
    for item_run in paragraph.xpath("./w:r[w:t]", namespaces=NS)[:6]:
        item_rpr = item_run.find("w:rPr", namespaces=NS)
        first_runs.append(
            {
                "text": xml_text(item_run)[:80],
                "rFonts": child_attrs(item_rpr, "rFonts"),
                "sz": child_attrs(item_rpr, "sz"),
                "szCs": child_attrs(item_rpr, "szCs"),
                "spacing": child_attrs(item_rpr, "spacing"),
                "b": child_attrs(item_rpr, "b"),
                "i": child_attrs(item_rpr, "i"),
                "color": child_attrs(item_rpr, "color"),
            }
        )
    return {
        "index": index,
        "style_guess": guess_strict_style(xml_text(paragraph)),
        "text": xml_text(paragraph)[:160],
        "pPr": {
            "pStyle": child_attrs(ppr, "pStyle"),
            "jc": child_attrs(ppr, "jc"),
            "spacing": child_attrs(ppr, "spacing"),
            "ind": child_attrs(ppr, "ind"),
            "keepNext": child_attrs(ppr, "keepNext"),
        },
        "first_run_rPr": {
            "rStyle": child_attrs(rpr, "rStyle"),
            "rFonts": child_attrs(rpr, "rFonts"),
            "sz": child_attrs(rpr, "sz"),
            "szCs": child_attrs(rpr, "szCs"),
            "spacing": child_attrs(rpr, "spacing"),
            "b": child_attrs(rpr, "b"),
            "i": child_attrs(rpr, "i"),
            "vertAlign": child_attrs(rpr, "vertAlign"),
            "color": child_attrs(rpr, "color"),
        },
        "first_runs": first_runs,
    }

def normal_style_observation(styles_root) -> dict[str, Any]:
    if styles_root is None:
        return {}
    matches = styles_root.xpath(".//w:style[@w:type='paragraph' and @w:styleId='Normal']", namespaces=NS)
    if not matches:
        matches = styles_root.xpath(".//w:style[@w:type='paragraph' and @w:default='1']", namespaces=NS)
    style_el = matches[0] if matches else None
    ppr = style_el.find("w:pPr", namespaces=NS) if style_el is not None else None
    rpr = style_el.find("w:rPr", namespaces=NS) if style_el is not None else None
    defaults_ppr = styles_root.find(".//w:docDefaults/w:pPrDefault/w:pPr", namespaces=NS)
    defaults_rpr = styles_root.find(".//w:docDefaults/w:rPrDefault/w:rPr", namespaces=NS)
    return {
        "style": {
            "styleId": local_attrs(style_el).get("styleId") if style_el is not None else None,
            "jc": child_attrs(ppr, "jc"),
            "spacing": child_attrs(ppr, "spacing"),
            "ind": child_attrs(ppr, "ind"),
            "rFonts": child_attrs(rpr, "rFonts"),
            "sz": child_attrs(rpr, "sz"),
            "szCs": child_attrs(rpr, "szCs"),
        },
        "docDefaults": {
            "spacing": child_attrs(defaults_ppr, "spacing"),
            "ind": child_attrs(defaults_ppr, "ind"),
            "rFonts": child_attrs(defaults_rpr, "rFonts"),
            "sz": child_attrs(defaults_rpr, "sz"),
            "szCs": child_attrs(defaults_rpr, "szCs"),
        },
    }

def strict_ooxml_observations(document_root, styles_root) -> dict[str, Any]:
    sect_prs = document_root.xpath(".//w:sectPr", namespaces=NS) if document_root is not None else []
    sect_pr = sect_prs[-1] if sect_prs else None
    observations: dict[str, Any] = {
        "page_ooxml": {
            "pgSz": child_attrs(sect_pr, "pgSz"),
            "pgMar": child_attrs(sect_pr, "pgMar"),
            "docGrid": child_attrs(sect_pr, "docGrid"),
        },
        "normal_style": normal_style_observation(styles_root),
        "paragraphs": [],
    }
    seen: set[str] = set()
    paragraphs = document_root.xpath(".//w:body/w:p", namespaces=NS) if document_root is not None else []
    for index, paragraph in enumerate(paragraphs):
        style_name = guess_strict_style(xml_text(paragraph))
        if not style_name or style_name in seen:
            continue
        observations["paragraphs"].append(raw_paragraph_observation(index, paragraph))
        seen.add(style_name)
    return observations

def cell_border_info(cell):
    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return {}
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        return {}
    info = {}
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{edge}"))
        if element is not None:
            info[edge] = attr_map(element)
    return info

def table_info(i, table):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW")) if tbl_pr is not None else None
    tbl_jc = tbl_pr.find(qn("w:jc")) if tbl_pr is not None else None
    tbl_layout = tbl_pr.find(qn("w:tblLayout")) if tbl_pr is not None else None
    tbl_borders = tbl_pr.find(qn("w:tblBorders")) if tbl_pr is not None else None
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar")) if tbl_pr is not None else None
    margins = {}
    if tbl_cell_mar is not None:
        for edge in ["top", "left", "bottom", "right"]:
            margins[edge] = attr_map(tbl_cell_mar.find(qn(f"w:{edge}")))

    rows = []
    for row in table.rows[:4]:
        rows.append([cell.text[:80] for cell in row.cells])

    sampled_cell_borders = []
    sample_rows = []
    if table.rows:
        sample_rows.append(("first", table.rows[0]))
        if len(table.rows) > 1:
            sample_rows.append(("last", table.rows[-1]))
    for label, row in sample_rows:
        sampled_cell_borders.append(
            {
                "row": label,
                "cells": [cell_border_info(cell) for cell in row.cells[:8]],
            }
        )

    return {
        "index": i,
        "style": table.style.name if table.style else None,
        "rows": len(table.rows),
        "cols": len(table.columns) if table.rows else 0,
        "tbl_width": attr_map(tbl_w),
        "alignment": attr_map(tbl_jc),
        "layout": attr_map(tbl_layout),
        "table_borders": attr_map(tbl_borders),
        "cell_margins": margins,
        "sample_rows": rows,
        "sampled_cell_borders": sampled_cell_borders,
    }

def image_layouts_from_xml(document_xml: str):
    layouts = []
    for match in re.finditer(r"<wp:(inline|anchor)\b[\s\S]*?</wp:\1>", document_xml):
        kind = match.group(1)
        fragment = match.group(0)
        extent = re.search(r"<wp:extent\b[^>]*cx=\"(\d+)\"[^>]*cy=\"(\d+)\"", fragment)
        doc_pr = re.search(r"<wp:docPr\b[^>]*id=\"([^\"]+)\"[^>]*name=\"([^\"]*)\"[^>]*", fragment)
        distances = {}
        for key in ["distT", "distB", "distL", "distR"]:
            dist = re.search(rf"\b{key}=\"([^\"]+)\"", fragment)
            distances[key] = int(dist.group(1)) if dist else None
        item = {
            "kind": kind,
            "docPr_id": doc_pr.group(1) if doc_pr else None,
            "name": doc_pr.group(2) if doc_pr else None,
            "distances": distances,
        }
        if extent:
            cx, cy = int(extent.group(1)), int(extent.group(2))
            item["width_mm"] = round(cx / 36000, 2)
            item["height_mm"] = round(cy / 36000, 2)
        layouts.append(item)
    return layouts

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx", type=Path)
    ap.add_argument("--out", type=Path, default=Path("extracted_style_map.json"))
    args = ap.parse_args()

    doc = Document(args.docx)
    sec = doc.sections[0]
    report: dict[str, Any] = {
        "source": str(args.docx),
        "section": {
            "page_width_cm": cm(sec.page_width),
            "page_height_cm": cm(sec.page_height),
            "top_margin_cm": cm(sec.top_margin),
            "bottom_margin_cm": cm(sec.bottom_margin),
            "left_margin_cm": cm(sec.left_margin),
            "right_margin_cm": cm(sec.right_margin),
            "header_distance_cm": cm(sec.header_distance),
            "footer_distance_cm": cm(sec.footer_distance),
        },
        "counts": {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "inline_shapes": len(doc.inline_shapes)},
        "paragraph_samples": [],
        "all_styles_used": sorted({p.style.name for p in doc.paragraphs if p.style}),
        "tables": [],
        "image_layouts": [],
        "ooxml_features": {},
        "strict_ooxml": {},
    }

    patterns = [
        r"页面设置", r"文章类型", r"^DOI", r"摘要：", r"关键词：", r"中图分类号",
        r"^Analysis", r"^Abstract:", r"^Keywords:", r"^\d+(?:\.\d+)*\s+", r"^图\s*\d+", r"^表\s*\d+", r"^\(\d+\)$", r"参考文献", r"符号说明"
    ]
    combined = re.compile("|".join(patterns))
    for i, p in enumerate(doc.paragraphs):
        if combined.search(p.text.strip()):
            report["paragraph_samples"].append(paragraph_info(i, p))

    for ti, table in enumerate(doc.tables):
        report["tables"].append(table_info(ti, table))

    with zipfile.ZipFile(args.docx) as z:
        names = set(z.namelist())
        document_xml = z.read("word/document.xml").decode("utf-8", errors="ignore") if "word/document.xml" in names else ""
        document_root = etree.fromstring(document_xml.encode("utf-8")) if document_xml else None
        styles_root = etree.fromstring(z.read("word/styles.xml")) if "word/styles.xml" in names else None
        report["strict_ooxml"] = strict_ooxml_observations(document_root, styles_root)
        report["image_layouts"] = image_layouts_from_xml(document_xml)
        report["ooxml_features"] = {
            "has_hyperlinks": "<w:hyperlink" in document_xml or any("/hyperlink" in n for n in names),
            "has_fields": "w:fldChar" in document_xml or "w:instrText" in document_xml,
            "has_textboxes": "w:txbxContent" in document_xml or "v:textbox" in document_xml,
            "has_comments_part": "word/comments.xml" in names,
            "has_footnotes_part": "word/footnotes.xml" in names,
            "has_endnotes_part": "word/endnotes.xml" in names,
            "has_headers": any(n.startswith("word/header") for n in names),
            "has_footers": any(n.startswith("word/footer") for n in names),
            "section_properties_count": document_xml.count("<w:sectPr"),
        }

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {args.out}")

if __name__ == "__main__":
    main()
